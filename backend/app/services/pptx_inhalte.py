from __future__ import annotations

from app.core.fehler import leise

MAX_PUNKTE = 7
MAX_ZEILEN = 8
MAX_SPALTEN = 6


def zeilen_aus(daten) -> list[list[str]]:
    if isinstance(daten, dict):
        kopf = list(daten.keys())
        werte = list(daten.values())
        if werte and isinstance(werte[0], (list, tuple)):
            zeilen = [kopf]
            for stelle in range(max(len(w) for w in werte)):
                zeilen.append(
                    [
                        str(w[stelle]) if stelle < len(w) else ""
                        for w in werte
                    ]
                )
            return zeilen
        return [kopf, [str(w) for w in werte]]
    if not isinstance(daten, (list, tuple)):
        return []
    zeilen: list[list[str]] = []
    for eintrag in daten:
        if isinstance(eintrag, (list, tuple)):
            zeilen.append([str(z) for z in eintrag])
        elif isinstance(eintrag, dict):
            if not zeilen:
                zeilen.append([str(k) for k in eintrag.keys()])
            zeilen.append([str(w) for w in eintrag.values()])
        else:
            zeilen.append([str(eintrag)])
    return zeilen


def punkt_text(eintrag) -> tuple[str, str]:
    if isinstance(eintrag, dict):
        titel = str(eintrag.get("titel") or eintrag.get("title") or "").strip()
        text = str(eintrag.get("text") or eintrag.get("beschreibung") or "").strip()
        if titel and text:
            return titel, text
        return (titel or text), ""
    roh = str(eintrag or "").strip()
    if ":" in roh and len(roh.split(":", 1)[0]) <= 34:
        kopf, rest = roh.split(":", 1)
        return kopf.strip(), rest.strip()
    return roh, ""


def tabelle(slide, daten, left: float, top: float, width: float, height: float,
            colors: dict[str, str], rgb, schrift: str) -> object | None:
    from pptx.util import Inches, Pt

    zeilen = zeilen_aus(daten)[:MAX_ZEILEN]
    if not zeilen:
        return None
    spalten = min(MAX_SPALTEN, max(len(z) for z in zeilen))
    if spalten < 1:
        return None
    try:
        form = slide.shapes.add_table(
            len(zeilen), spalten, Inches(left), Inches(top), Inches(width),
            Inches(min(height, 0.52 * len(zeilen))),
        )
    except Exception as fehler:
        leise(fehler, "services/pptx_inhalte")
        return None
    tab = form.table
    tab.first_row = True
    for r, zeile in enumerate(zeilen):
        tab.rows[r].height = Inches(0.5)
        for c in range(spalten):
            zelle = tab.cell(r, c)
            zelle.text = zeile[c] if c < len(zeile) else ""
            zelle.margin_left = Inches(0.14)
            zelle.margin_right = Inches(0.1)
            zelle.fill.solid()
            zelle.fill.fore_color.rgb = rgb(
                colors["primary"] if r == 0 else (
                    colors["light"] if r % 2 else colors["soft"]
                )
            )
            for absatz in zelle.text_frame.paragraphs:
                absatz.space_after = Pt(0)
                for lauf in absatz.runs:
                    lauf.font.size = Pt(15 if r == 0 else 14)
                    lauf.font.bold = r == 0
                    lauf.font.name = schrift
                    lauf.font.color.rgb = rgb(
                        colors["light"] if r == 0 else colors["dark"]
                    )
    return form


def diagramm(slide, daten: dict, left: float, top: float, width: float, height: float,
             colors: dict[str, str], rgb, schrift: str) -> object | None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    arten = {
        "balken": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "saeulen": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "linie": XL_CHART_TYPE.LINE_MARKERS,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "flaeche": XL_CHART_TYPE.AREA,
        "kreis": XL_CHART_TYPE.PIE,
        "pie": XL_CHART_TYPE.PIE,
        "donut": XL_CHART_TYPE.DOUGHNUT,
        "gestapelt": XL_CHART_TYPE.COLUMN_STACKED,
    }
    art = arten.get(str(daten.get("art", "balken")).strip().lower(),
                    XL_CHART_TYPE.COLUMN_CLUSTERED)
    kategorien = [str(k) for k in (daten.get("kategorien") or daten.get("labels") or [])]
    reihen = daten.get("reihen") or daten.get("series") or []
    if isinstance(reihen, dict):
        reihen = [{"name": k, "werte": v} for k, v in reihen.items()]
    sauber: list[tuple[str, list[float]]] = []
    for reihe in reihen:
        if not isinstance(reihe, dict):
            continue
        name = str(reihe.get("name") or reihe.get("titel") or "Werte")
        werte = []
        for w in reihe.get("werte") or reihe.get("values") or []:
            try:
                werte.append(float(str(w).replace(",", ".").replace("%", "").strip()))
            except (TypeError, ValueError):
                werte.append(0.0)
        if werte:
            sauber.append((name, werte))
    if not kategorien or not sauber:
        return None
    laenge = len(kategorien)
    inhalt = CategoryChartData()
    inhalt.categories = kategorien
    for name, werte in sauber[:4]:
        gefuellt = (werte + [0.0] * laenge)[:laenge]
        inhalt.add_series(name, gefuellt)
    try:
        rahmen = slide.shapes.add_chart(
            art, Inches(left), Inches(top), Inches(width), Inches(height), inhalt
        )
    except Exception as fehler:
        leise(fehler, "services/pptx_inhalte")
        return None
    diagramm_ = rahmen.chart
    diagramm_.font.size = Pt(13)
    diagramm_.font.name = schrift
    diagramm_.font.color.rgb = rgb(colors["muted"])
    if len(sauber) > 1 or art in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
        diagramm_.has_legend = True
        diagramm_.legend.position = XL_LEGEND_POSITION.BOTTOM
        diagramm_.legend.include_in_layout = False
    else:
        diagramm_.has_legend = False
    paletten = [colors["primary"], colors["accent"], colors["secondary"], colors["muted"]]
    try:
        if art in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
            punkte = diagramm_.plots[0].series[0].points
            for stelle, punkt in enumerate(punkte):
                punkt.format.fill.solid()
                punkt.format.fill.fore_color.rgb = rgb(paletten[stelle % len(paletten)])
        else:
            for stelle, reihe in enumerate(diagramm_.series):
                reihe.format.fill.solid()
                reihe.format.fill.fore_color.rgb = rgb(paletten[stelle % len(paletten)])
    except Exception as fehler:
        leise(fehler, "services/pptx_inhalte")
    try:
        diagramm_.plots[0].has_data_labels = art not in (XL_CHART_TYPE.LINE_MARKERS,)
        if diagramm_.plots[0].has_data_labels:
            beschriftung = diagramm_.plots[0].data_labels
            beschriftung.font.size = Pt(12)
            beschriftung.font.bold = True
            beschriftung.font.name = schrift
            beschriftung.font.color.rgb = rgb(colors["dark"])
    except Exception as fehler:
        leise(fehler, "services/pptx_inhalte")
    return rahmen


def absaetze(slide, texte: list[str], left: float, top: float, width: float,
             height: float, size: int, colors: dict[str, str], rgb, schrift: str):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    rahmen = box.text_frame
    rahmen.word_wrap = True
    rahmen.margin_left = 0
    rahmen.margin_right = 0
    for stelle, roh in enumerate(texte):
        absatz = rahmen.paragraphs[0] if stelle == 0 else rahmen.add_paragraph()
        absatz.space_after = Pt(12)
        absatz.line_spacing = 1.22
        kopf, rest = punkt_text(roh)
        if rest:
            fett = absatz.add_run()
            fett.text = kopf + " "
            fett.font.bold = True
            fett.font.size = Pt(size)
            fett.font.name = schrift
            fett.font.color.rgb = rgb(colors["primary"])
            lauf = absatz.add_run()
            lauf.text = rest
        else:
            lauf = absatz.add_run()
            lauf.text = kopf
        lauf.font.size = Pt(size)
        lauf.font.name = schrift
        lauf.font.color.rgb = rgb(colors["muted"])
    return box


BILD_STIL = (
    "clean modern editorial illustration, flat vector style, generous negative space, "
    "muted professional palette, no text, no words, no letters, no watermark"
)


async def bild_besorgen(beschreibung: str, breit: bool = True) -> str:
    from app.services.studio_service import get_studio_service

    auftrag = str(beschreibung or "").strip()
    if not auftrag:
        return ""
    try:
        dienst = get_studio_service()
        werk = await dienst.generate(
            f"{auftrag}. {BILD_STIL}",
            "bild",
            "",
            "1280x720" if breit else "1024x1024",
            "text, words, letters, watermark, logo, signature",
            "",
            "",
        )
        return str(dienst.file(str(werk["datei"])))
    except Exception as fehler:
        leise(fehler, "services/pptx_inhalte")
        return ""
