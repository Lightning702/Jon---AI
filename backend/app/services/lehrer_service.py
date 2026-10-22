from __future__ import annotations

import asyncio
import json
import random
import re
import threading
import time
import uuid
from datetime import date, datetime
from pathlib import Path

from app.core.config import DATA_DIR
from app.core.fehler import leise
from app.core.store import atomic_write_text

ORDNER = DATA_DIR / "lehrer"
STUECK = 22_000
MATERIAL_MAX = 600_000
INTERVALLE = [120, 600, 86400, 3 * 86400, 7 * 86400, 16 * 86400, 35 * 86400]
ARTEN = ("auswahl", "frei", "lueckentext", "wahr_falsch")

LEHRER_ROLLE = """Du bist Jon Lehrer, ein geduldiger, begeisternder und fachlich exzellenter Privatlehrer.
Du unterrichtest auf Deutsch (ausser die Lernsprache verlangt etwas anderes), klar, freundlich und auf Augenhöhe des Lernenden.
Du stützt dich zuerst auf das Lernmaterial des Lernenden. Allgemeinwissen ergänzt du nur, wenn das Material etwas nicht abdeckt, und bleibst dabei sachlich korrekt.
Erfinde keine Inhalte, die dem Material widersprechen. Rechne und prüfe alles nach, bevor du es behauptest."""

GLIEDERUNG_AUFTRAG = """Plane aus dem Lernstoff einen Kurs.
Fach/Thema: {fach}
Ziel: {ziel}
Niveau: {niveau}
{termin}
Antworte NUR mit JSON: {{"titel": "kurzer Kurstitel", "zusammenfassung": "3-5 Sätze, worum es geht und was man am Ende kann", "lektionen": [{{"titel": "...", "ziele": ["Ich kann ...", "..."], "stoff": "welche Regeln, Begriffe, Vokabeln oder Inhalte aus dem Material hierher gehören"}}]}}
4 bis 10 Lektionen in sinnvoller Reihenfolge vom Einfachen zum Schweren; die letzte darf eine gemischte Wiederholung sein. Alles muss zum Material passen.{sprache}"""

LEKTION_AUFTRAG = """Schreibe Lektion {nr} von {anzahl} des Kurses „{kurs}“ ({fach}, Niveau: {niveau}).
Lektion: {titel}
Lernziele: {ziele}
Stoff: {stoff}
Antworte NUR mit JSON: {{"inhalt": "verständliche Erklärung mit Merksätzen und Beispielen, 180-400 Wörter, Absätze mit \\n\\n, Merksätze mit **fett**", "beispiele": ["konkretes Beispiel"], "begriffe": ["Fachbegriff: kurze Erklärung"],
"aufgaben": [{{"art": "auswahl|frei|lueckentext|wahr_falsch", "frage": "...", "optionen": ["nur bei auswahl, 3-4 Antworten"], "loesung": "richtige Antwort (bei auswahl exakt der Optionstext, bei wahr_falsch 'wahr' oder 'falsch', bei lueckentext das fehlende Wort)", "erklaerung": "warum", "schwierigkeit": 1}}]}}
5 abwechslungsreiche Aufgaben mit steigender Schwierigkeit (1-3), Lücken mit ___ markieren. Halte dich an das Material.{sprache}"""

PLAN_AUFTRAG = """Erstelle aus dem Lernstoff einen kompletten Kurs.
Fach/Thema: {fach}
Ziel: {ziel}
Niveau: {niveau}
{termin}
Antworte NUR mit JSON in genau diesem Format:
{{"titel": "kurzer Kurstitel", "zusammenfassung": "3-5 Sätze, worum es geht und was man am Ende kann",
"lektionen": [{{"titel": "...", "ziele": ["Ich kann ..."], "inhalt": "verständliche Erklärung mit Merksätzen und Beispielen, 180-400 Wörter, Absätze mit \\n\\n, Merksätze mit **fett**", "beispiele": ["konkretes Beispiel"], "begriffe": ["Fachbegriff: kurze Erklärung"],
"aufgaben": [{{"art": "auswahl|frei|lueckentext|wahr_falsch", "frage": "...", "optionen": ["nur bei auswahl, 3-4 Antworten"], "loesung": "die richtige Antwort (bei auswahl der exakte Optionstext, bei wahr_falsch 'wahr' oder 'falsch', bei lueckentext das fehlende Wort)", "erklaerung": "warum das richtig ist", "schwierigkeit": 1}}]}}]}}
Regeln: 4 bis 10 Lektionen in sinnvoller Reihenfolge vom Einfachen zum Schweren, je Lektion 4 bis 6 abwechslungsreiche Aufgaben (Schwierigkeit 1-3). Lückentext-Fragen markieren die Lücke mit ___. Alles muss zum Material passen.{sprache}"""

THEMEN_AUFTRAG = """Lies diesen Ausschnitt aus Lernmaterial und liste die Lerninhalte auf. Antworte NUR mit JSON:
{"themen": [{"titel": "...", "kern": ["wichtige Aussage, Regel, Formel, Vokabel oder Definition", "..."]}]}
Höchstens 12 Themen, nichts erfinden, Fachbegriffe exakt übernehmen."""

AUFGABEN_AUFTRAG = """Erstelle {anzahl} neue, abwechslungsreiche Übungsaufgaben zur Lektion „{titel}“.
Lerninhalt der Lektion:
{inhalt}
Bereits gestellt (nicht wiederholen): {bisher}
Schwerpunkt: {schwerpunkt}
Antworte NUR mit JSON: {{"aufgaben": [{{"art": "auswahl|frei|lueckentext|wahr_falsch", "frage": "...", "optionen": ["nur bei auswahl"], "loesung": "...", "erklaerung": "...", "schwierigkeit": 1}}]}}"""

BEWERTEN_AUFTRAG = """Bewerte die Antwort eines Lernenden fair, aber genau.
Aufgabe: {frage}
Musterlösung: {loesung}
Erklärung: {erklaerung}
Antwort des Lernenden: {antwort}
Kleine Tippfehler und gleichwertige Formulierungen zählen als richtig, inhaltliche Fehler nicht. Bei Sprach- oder Rechtschreibaufgaben zählt die exakte Form.
Antworte NUR mit JSON: {{"richtig": true, "teilweise": false, "feedback": "1-3 freundliche Sätze: was gut war, was falsch war und wie man es sich merkt"}}"""

SPRACHFAECHER = ("deutsch", "englisch", "english", "französisch", "franzoesisch", "spanisch", "italienisch", "latein", "russisch", "türkisch", "tuerkisch", "japanisch", "chinesisch", "niederländisch", "portugiesisch", "polnisch", "kroatisch", "daf", "sprache", "vokabel", "grammatik")


def _jetzt() -> float:
    return time.time()


def _json_aus(text: str) -> dict:
    roh = (text or "").strip()
    roh = re.sub(r"^```(?:json)?\s*|\s*```$", "", roh)
    dekoder = json.JSONDecoder()
    bestes: dict = {}
    position = 0
    while True:
        start = roh.find("{", position)
        if start < 0:
            return bestes
        try:
            wert, ende = dekoder.raw_decode(roh, start)
        except ValueError:
            position = start + 1
            continue
        if isinstance(wert, dict) and len(json.dumps(wert)) > len(json.dumps(bestes)):
            bestes = wert
        position = ende


def _sauber(text, grenze: int = 4000) -> str:
    return re.sub(r"[ \t]+", " ", str(text or "")).strip()[:grenze]


def _aufgabe(roh: dict, lektion: int) -> dict | None:
    if not isinstance(roh, dict):
        return None
    frage = _sauber(roh.get("frage"), 1200)
    loesung = _sauber(roh.get("loesung"), 1200)
    if not frage or not loesung:
        return None
    art = str(roh.get("art") or "frei").lower().replace("-", "_").replace("ü", "ue")
    if art not in ARTEN:
        art = "frei"
    optionen = [_sauber(o, 300) for o in (roh.get("optionen") or []) if _sauber(o, 300)][:6]
    if art == "auswahl":
        if len(optionen) < 2:
            art = "frei"
            optionen = []
        elif not any(_gleich(o, loesung) for o in optionen):
            optionen = optionen[:3] + [loesung]
        random.shuffle(optionen)
    else:
        optionen = []
    if art == "wahr_falsch":
        loesung = "wahr" if _gleich(loesung, "wahr") or _gleich(loesung, "richtig") or _gleich(loesung, "true") else "falsch"
    try:
        schwierigkeit = max(1, min(3, int(roh.get("schwierigkeit") or 1)))
    except (TypeError, ValueError):
        schwierigkeit = 1
    return {
        "id": uuid.uuid4().hex[:12],
        "lektion": lektion,
        "art": art,
        "frage": frage,
        "optionen": optionen,
        "loesung": loesung,
        "erklaerung": _sauber(roh.get("erklaerung"), 1500),
        "schwierigkeit": schwierigkeit,
        "stufe": 0,
        "faellig": 0.0,
        "richtig": 0,
        "falsch": 0,
        "gesehen": 0,
    }


def _normal(text: str) -> str:
    klein = str(text or "").casefold().strip()
    klein = re.sub(r"[\s.,;:!?\"'„“()]+", " ", klein)
    return klein.strip()


def _gleich(a: str, b: str) -> bool:
    return _normal(a) == _normal(b)


def _sprachfach(fach: str) -> bool:
    klein = (fach or "").casefold()
    return any(wort in klein for wort in SPRACHFAECHER)


class LehrerService:
    def __init__(self, ordner: Path | None = None) -> None:
        self.ordner = Path(ordner or ORDNER)
        self.lock = threading.RLock()
        self.laeufe: dict[str, asyncio.Task] = {}

    def _pfad(self, kennung: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{12}", kennung or ""):
            raise ValueError("Unbekannter Kurs.")
        return self.ordner / kennung

    def laden(self, kennung: str) -> dict:
        try:
            return json.loads((self._pfad(kennung) / "kurs.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            raise ValueError("Diesen Kurs gibt es nicht.")

    def speichern(self, kurs: dict) -> None:
        with self.lock:
            ordner = self._pfad(kurs["id"])
            ordner.mkdir(parents=True, exist_ok=True)
            atomic_write_text(ordner / "kurs.json", json.dumps(kurs, ensure_ascii=False, indent=1))

    def material(self, kennung: str) -> str:
        try:
            return (self._pfad(kennung) / "material.txt").read_text(encoding="utf-8")
        except OSError:
            return ""

    def liste(self) -> list[dict]:
        kurse = []
        if not self.ordner.exists():
            return kurse
        for datei in self.ordner.glob("*/kurs.json"):
            try:
                kurs = json.loads(datei.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            kurse.append(self.kurz(kurs))
        kurse.sort(key=lambda k: k.get("zuletzt", 0), reverse=True)
        return kurse

    def kurz(self, kurs: dict) -> dict:
        lektionen = kurs.get("lektionen") or []
        beherrschung = round(sum(l.get("beherrschung", 0) for l in lektionen) / len(lektionen), 2) if lektionen else 0.0
        faellig = sum(1 for a in kurs.get("aufgaben", []) if a.get("gesehen") and a.get("faellig", 0) <= _jetzt())
        return {
            "id": kurs["id"],
            "titel": kurs.get("titel") or kurs.get("fach"),
            "fach": kurs.get("fach", ""),
            "ziel": kurs.get("ziel", ""),
            "termin": kurs.get("termin", ""),
            "status": kurs.get("status", ""),
            "fortschritt": kurs.get("fortschritt", 0),
            "schritt": kurs.get("schritt", ""),
            "fehler": kurs.get("fehler", ""),
            "lektionen": len(lektionen),
            "beherrschung": beherrschung,
            "faellig": faellig,
            "dateien": len(kurs.get("dateien", [])),
            "zuletzt": kurs.get("zuletzt", kurs.get("erstellt", 0)),
            "serie": kurs.get("serie", 0),
        }

    def ansicht(self, kennung: str) -> dict:
        kurs = self.laden(kennung)
        ergebnis = {**self.kurz(kurs), "zusammenfassung": kurs.get("zusammenfassung", ""), "niveau": kurs.get("niveau", ""), "dateiliste": [{"name": d.get("name"), "zeichen": d.get("zeichen", 0)} for d in kurs.get("dateien", [])]}
        ergebnis["lektionsliste"] = [
            {
                "nr": l["nr"],
                "titel": l.get("titel", ""),
                "ziele": l.get("ziele", []),
                "inhalt": l.get("inhalt", ""),
                "beispiele": l.get("beispiele", []),
                "begriffe": l.get("begriffe", []),
                "beherrschung": l.get("beherrschung", 0),
                "bereit": l.get("bereit", True),
                "aufgaben": sum(1 for a in kurs.get("aufgaben", []) if a.get("lektion") == l["nr"]),
            }
            for l in kurs.get("lektionen", [])
        ]
        ergebnis["chat"] = kurs.get("chat", [])[-30:]
        ergebnis["empfehlung"] = self.empfehlung(kurs)
        return ergebnis

    def empfehlung(self, kurs: dict) -> str:
        if kurs.get("status") != "bereit":
            return ""
        lektionen = kurs.get("lektionen") or []
        faellig = sum(1 for a in kurs.get("aufgaben", []) if a.get("gesehen") and a.get("faellig", 0) <= _jetzt())
        tage = None
        if kurs.get("termin"):
            try:
                tage = (date.fromisoformat(kurs["termin"]) - date.today()).days
            except ValueError:
                tage = None
        schwach = sorted(lektionen, key=lambda l: l.get("beherrschung", 0))
        teile = []
        if tage is not None and tage >= 0:
            teile.append(f"Noch {tage} Tage bis zum Termin." if tage else "Heute ist der Termin – kurz wiederholen, nicht mehr neu lernen.")
        if faellig:
            teile.append("1 Aufgabe ist zur Wiederholung fällig." if faellig == 1 else f"{faellig} Aufgaben sind zur Wiederholung fällig.")
        if schwach and schwach[0].get("beherrschung", 0) < 0.8:
            teile.append(f"Als Nächstes: „{schwach[0].get('titel')}“.")
        elif lektionen:
            teile.append("Alles sitzt gut – mach einen Probetest.")
        return " ".join(teile)

    def erstellen(self, fach: str, ziel: str = "", niveau: str = "", termin: str = "", dateien: list[str] | None = None) -> dict:
        fach = _sauber(fach, 160)
        if len(fach) < 2:
            raise ValueError("Sag mir, was du lernen möchtest – zum Beispiel „Deutsch: Kommasetzung“.")
        if termin:
            try:
                termin = date.fromisoformat(termin.strip()[:10]).isoformat()
            except ValueError:
                termin = ""
        kurs = {
            "id": uuid.uuid4().hex[:12],
            "titel": fach,
            "fach": fach,
            "ziel": _sauber(ziel, 300),
            "niveau": _sauber(niveau, 120),
            "termin": termin,
            "erstellt": _jetzt(),
            "zuletzt": _jetzt(),
            "status": "baut",
            "fortschritt": 0,
            "schritt": "Lerndateien werden gelesen",
            "dateien": [{"pfad": str(p)} for p in (dateien or []) if str(p).strip()][:40],
            "lektionen": [],
            "aufgaben": [],
            "verlauf": [],
            "chat": [],
            "serie": 0,
        }
        self.speichern(kurs)
        return kurs

    def starten(self, kurs: dict) -> None:
        kennung = kurs["id"]
        alt = self.laeufe.get(kennung)
        if alt and not alt.done():
            return
        self.laeufe[kennung] = asyncio.create_task(self.aufbauen(kennung))

    def _schritt(self, kennung: str, schritt: str, fortschritt: int) -> None:
        with self.lock:
            kurs = self.laden(kennung)
            kurs.update(schritt=schritt, fortschritt=fortschritt)
            self.speichern(kurs)

    async def _texte_lesen(self, kurs: dict) -> str:
        from app.services.attachment_service import get_attachment_service

        teile = []
        dienst = get_attachment_service()
        gelesen = []
        for eintrag in kurs.get("dateien", []):
            pfad = Path(str(eintrag.get("pfad") or "")).expanduser()
            if not pfad.is_file():
                gelesen.append({"pfad": str(pfad), "name": pfad.name, "zeichen": 0, "fehler": "nicht gefunden"})
                continue
            try:
                ergebnis = await dienst.auswerten(pfad, pfad.name, "", beschreiben=True)
                text = str(ergebnis.get("content") or "")
                if ergebnis.get("kind") in ("pdf", "praesentation", "dokument", "tabelle", "text") and len(text) >= 19_000:
                    text = await asyncio.to_thread(_voller_text, pfad, text)
            except Exception as fehler:
                leise(fehler, "lehrer/lesen")
                text = ""
            gelesen.append({"pfad": str(pfad), "name": pfad.name, "zeichen": len(text)})
            if text.strip():
                teile.append(f"### Datei: {pfad.name}\n{text.strip()}")
        kurs["dateien"] = gelesen
        return "\n\n".join(teile)[:MATERIAL_MAX]

    async def _modell(self, system: str, nutzer: str, max_tokens: int = 6000, temperatur: float = 0.4) -> str:
        from app.services.llm import complete
        from app.services.settings_service import get_settings_service

        provider, model = get_settings_service().selection()
        return await complete(system, nutzer, provider=provider or None, model=model or None, max_tokens=max_tokens, temperature=temperatur, denken=False)

    async def aufbauen(self, kennung: str) -> None:
        try:
            kurs = self.laden(kennung)
            material = await self._texte_lesen(kurs)
            ordner = self._pfad(kennung)
            ordner.mkdir(parents=True, exist_ok=True)
            atomic_write_text(ordner / "material.txt", material)
            with self.lock:
                frisch = self.laden(kennung)
                frisch["dateien"] = kurs["dateien"]
                self.speichern(frisch)
            self._schritt(kennung, "Jon liest den Stoff", 25)
            if len(material) > STUECK * 1.6:
                stuecke = [material[i : i + STUECK] for i in range(0, len(material), STUECK)][:24]
                grenze = asyncio.Semaphore(3)

                async def themen(stueck: str) -> str:
                    async with grenze:
                        try:
                            antwort = await self._modell(THEMEN_AUFTRAG, stueck, 2500, 0.2)
                        except Exception as fehler:
                            leise(fehler, "lehrer/themen")
                            return ""
                        daten = _json_aus(antwort)
                        zeilen = []
                        for thema in daten.get("themen") or []:
                            if isinstance(thema, dict):
                                zeilen.append(f"- {thema.get('titel')}: " + "; ".join(str(k) for k in thema.get("kern") or []))
                        return "\n".join(zeilen)

                ergebnisse = await asyncio.gather(*(themen(s) for s in stuecke))
                stoff = "ÜBERSICHT DES GESAMTEN MATERIALS:\n" + "\n".join(e for e in ergebnisse if e) + "\n\nAUSZUG AUS DEM MATERIAL:\n" + material[:STUECK]
            else:
                stoff = material or "(Kein Material hochgeladen – nutze dein Fachwissen zum Thema.)"
            self._schritt(kennung, "Jon plant die Lektionen", 40)
            sprache = "\nDas ist ein Sprachfach: baue Wortschatz, Grammatik mit Beispielsätzen, typische Fehler und kurze Übersetzungs- oder Schreibaufgaben ein." if _sprachfach(kurs.get("fach", "")) else ""
            auftrag = GLIEDERUNG_AUFTRAG.format(fach=kurs.get("fach"), ziel=kurs.get("ziel") or "den Stoff sicher beherrschen", niveau=kurs.get("niveau") or "passend zum Material", termin=f"Termin (Prüfung/Schularbeit): {kurs['termin']}" if kurs.get("termin") else "", sprache=sprache)
            plan: dict = {}
            for versuch in range(3):
                try:
                    plan = _json_aus(await self._modell(LEHRER_ROLLE, auftrag + "\n\nLERNSTOFF:\n" + stoff[:50_000], 3500, 0.3))
                except Exception as fehler:
                    if versuch == 2:
                        raise
                    leise(fehler, "lehrer/plan")
                    await asyncio.sleep(2)
                    continue
                if plan.get("lektionen"):
                    break
            rohe = [l for l in plan.get("lektionen") or [] if isinstance(l, dict) and l.get("titel")][:10]
            if not rohe:
                raise ValueError("Aus dem Material ließ sich kein Kurs bauen. Versuch es mit anderen Dateien oder einem genaueren Thema.")
            with self.lock:
                kurs = self.laden(kennung)
                kurs.update(
                    titel=_sauber(plan.get("titel"), 120) or kurs.get("fach"),
                    zusammenfassung=_sauber(plan.get("zusammenfassung"), 1500),
                    lektionen=[{"nr": nr, "titel": _sauber(r.get("titel"), 140), "ziele": [_sauber(z, 200) for z in r.get("ziele") or [] if _sauber(z, 200)][:6], "inhalt": "", "beispiele": [], "begriffe": [], "beherrschung": 0.0, "bereit": False} for nr, r in enumerate(rohe, 1)],
                    schritt=f"Lektion 1 von {len(rohe)} wird geschrieben",
                    fortschritt=50,
                )
                self.speichern(kurs)
            grenze = asyncio.Semaphore(8)
            fertig = 0

            async def lektion_bauen(nr: int, roh: dict) -> None:
                nonlocal fertig
                async with grenze:
                    ziele = "; ".join(str(z) for z in roh.get("ziele") or [])
                    auszug = self._auszug(material, f"{roh.get('titel')} {ziele} {roh.get('stoff', '')}") if len(material) > 14_000 else material
                    text = LEKTION_AUFTRAG.format(nr=nr, anzahl=len(rohe), kurs=plan.get("titel") or kurs.get("fach"), fach=kurs.get("fach"), niveau=kurs.get("niveau") or "passend zum Material", titel=roh.get("titel"), ziele=ziele, stoff=_sauber(roh.get("stoff"), 1500), sprache=sprache)
                    daten: dict = {}
                    for versuch in range(2):
                        try:
                            daten = _json_aus(await self._modell(LEHRER_ROLLE, text + "\n\nMATERIAL:\n" + (auszug or "(kein Material, nutze dein Fachwissen)")[:16_000], 5000, 0.4))
                        except Exception as fehler:
                            leise(fehler, "lehrer/lektion")
                            daten = {}
                        if daten.get("inhalt"):
                            break
                    aufgaben = [a for a in (_aufgabe(r, nr) for r in daten.get("aufgaben") or []) if a]
                    with self.lock:
                        aktuell = self.laden(kennung)
                        for eintrag in aktuell.get("lektionen", []):
                            if eintrag["nr"] == nr:
                                eintrag.update(
                                    inhalt=_sauber(daten.get("inhalt"), 6000).replace("\\n", "\n") or "Diese Lektion konnte gerade nicht geschrieben werden. Tippe auf „Neu bauen“.",
                                    beispiele=[_sauber(b, 500) for b in daten.get("beispiele") or [] if _sauber(b, 500)][:6],
                                    begriffe=[_sauber(b, 300) for b in daten.get("begriffe") or [] if _sauber(b, 300)][:12],
                                    bereit=True,
                                )
                        aktuell.setdefault("aufgaben", []).extend(aufgaben)
                        fertig += 1
                        aktuell["fortschritt"] = 50 + int(50 * fertig / len(rohe))
                        aktuell["schritt"] = "fertig" if fertig == len(rohe) else f"{fertig} von {len(rohe)} Lektionen fertig"
                        if any(e["nr"] == 1 and e.get("bereit") for e in aktuell["lektionen"]) and aktuell.get("aufgaben"):
                            aktuell["status"] = "bereit"
                        aktuell["zuletzt"] = _jetzt()
                        self.speichern(aktuell)

            await asyncio.gather(*(lektion_bauen(nr, roh) for nr, roh in enumerate(rohe, 1)))
            with self.lock:
                kurs = self.laden(kennung)
                if not kurs.get("aufgaben"):
                    raise ValueError("Die Übungen ließen sich gerade nicht erstellen. Tippe auf „Neu bauen“.")
                kurs.update(status="bereit", fortschritt=100, schritt="fertig", fehler="", zuletzt=_jetzt())
                self.speichern(kurs)
        except Exception as fehler:
            leise(fehler, "lehrer/aufbauen")
            try:
                with self.lock:
                    kurs = self.laden(kennung)
                    kurs.update(status="fehler", fehler=str(fehler)[:400] or type(fehler).__name__, schritt="Fehler")
                    self.speichern(kurs)
            except Exception as weiter:
                leise(weiter, "lehrer/aufbauen")

    async def warten(self, kennung: str, grenze: float = 600) -> dict:
        lauf = self.laeufe.get(kennung)
        if lauf is not None:
            try:
                await asyncio.wait_for(asyncio.shield(lauf), grenze)
            except asyncio.TimeoutError:
                pass
        return self.ansicht(kennung)

    def _oeffentlich(self, aufgabe: dict) -> dict:
        return {k: aufgabe[k] for k in ("id", "lektion", "art", "frage", "optionen", "schwierigkeit", "stufe")}

    async def naechste(self, kennung: str, lektion: int | None = None) -> dict:
        kurs = self.laden(kennung)
        if kurs.get("status") != "bereit":
            raise ValueError("Der Kurs ist noch nicht fertig.")
        jetzt = _jetzt()
        bereit = {l["nr"] for l in kurs.get("lektionen", []) if l.get("bereit", True)}
        alle = kurs.get("aufgaben", [])
        auswahl = [a for a in alle if lektion is None or a.get("lektion") == lektion]
        faellig = sorted((a for a in auswahl if a.get("gesehen") and a.get("faellig", 0) <= jetzt), key=lambda a: (a.get("stufe", 0), a.get("faellig", 0)))
        if faellig:
            return {"aufgabe": self._oeffentlich(faellig[0]), "grund": "Wiederholung"}
        neue = [a for a in auswahl if not a.get("gesehen")]
        if neue:
            reihenfolge = {l["nr"]: l.get("beherrschung", 0) for l in kurs.get("lektionen", [])}
            neue.sort(key=lambda a: (a.get("lektion", 0) if lektion is None else 0, a.get("schwierigkeit", 1), reihenfolge.get(a.get("lektion"), 0)))
            return {"aufgabe": self._oeffentlich(neue[0]), "grund": "neu"}
        kandidaten = [l for l in kurs.get("lektionen", []) if l["nr"] in bereit] or kurs.get("lektionen", [])
        ziel = lektion or min(kandidaten, key=lambda l: l.get("beherrschung", 0))["nr"]
        frisch = await self.mehr_aufgaben(kennung, ziel)
        if frisch:
            return {"aufgabe": self._oeffentlich(frisch[0]), "grund": "neu erstellt"}
        bald = sorted(auswahl, key=lambda a: a.get("faellig", 0))
        if bald:
            return {"aufgabe": self._oeffentlich(bald[0]), "grund": "vorgezogen"}
        raise ValueError("Gerade gibt es keine Aufgabe.")

    async def mehr_aufgaben(self, kennung: str, lektion: int, anzahl: int = 5, schwerpunkt: str = "") -> list[dict]:
        kurs = self.laden(kennung)
        eintrag = next((l for l in kurs.get("lektionen", []) if l["nr"] == lektion), None)
        if eintrag is None:
            return []
        bisher = [a["frage"][:120] for a in kurs.get("aufgaben", []) if a.get("lektion") == lektion][-25:]
        fehler = [a["frage"][:120] for a in kurs.get("aufgaben", []) if a.get("lektion") == lektion and a.get("falsch")]
        auftrag = AUFGABEN_AUFTRAG.format(anzahl=anzahl, titel=eintrag["titel"], inhalt=eintrag.get("inhalt", "")[:5000] + "\n" + "\n".join(eintrag.get("begriffe", [])), bisher=json.dumps(bisher, ensure_ascii=False), schwerpunkt=schwerpunkt or ("Übe gezielt das, was zuletzt falsch war: " + "; ".join(fehler[-5:]) if fehler else "abwechslungsreich, steigende Schwierigkeit"))
        try:
            antwort = await self._modell(LEHRER_ROLLE, auftrag, 5000, 0.6)
        except Exception as fehler_neu:
            leise(fehler_neu, "lehrer/aufgaben")
            return []
        neu = [a for a in (_aufgabe(r, lektion) for r in _json_aus(antwort).get("aufgaben") or []) if a]
        if neu:
            with self.lock:
                kurs = self.laden(kennung)
                kurs.setdefault("aufgaben", []).extend(neu)
                self.speichern(kurs)
        return neu

    async def antworten(self, kennung: str, aufgabe_id: str, antwort: str) -> dict:
        kurs = self.laden(kennung)
        aufgabe = next((a for a in kurs.get("aufgaben", []) if a["id"] == aufgabe_id), None)
        if aufgabe is None:
            raise ValueError("Diese Aufgabe gibt es nicht mehr.")
        antwort = str(antwort or "").strip()
        teilweise = False
        if not antwort:
            richtig = False
            feedback = "Keine Antwort ist auch eine Antwort – schau dir die Lösung an, dann klappt es beim nächsten Mal."
        elif aufgabe["art"] in ("auswahl", "wahr_falsch"):
            wahl = antwort
            if aufgabe["art"] == "wahr_falsch":
                wahl = "wahr" if _normal(antwort) in ("wahr", "richtig", "ja", "true", "w", "r") else "falsch"
            richtig = _gleich(wahl, aufgabe["loesung"])
            feedback = "Richtig!" if richtig else "Leider falsch."
        elif aufgabe["art"] == "lueckentext" and _gleich(antwort, aufgabe["loesung"]):
            richtig = True
            feedback = "Genau richtig!"
        else:
            try:
                bewertung = _json_aus(await self._modell(LEHRER_ROLLE, BEWERTEN_AUFTRAG.format(frage=aufgabe["frage"], loesung=aufgabe["loesung"], erklaerung=aufgabe.get("erklaerung", ""), antwort=antwort[:3000]), 900, 0.1))
                richtig = bool(bewertung.get("richtig"))
                teilweise = bool(bewertung.get("teilweise")) and not richtig
                feedback = _sauber(bewertung.get("feedback"), 1200) or ("Richtig!" if richtig else "Noch nicht ganz.")
            except Exception as fehler:
                leise(fehler, "lehrer/bewerten")
                richtig = _gleich(antwort, aufgabe["loesung"])
                feedback = "Richtig!" if richtig else "Das stimmt so noch nicht."
        jetzt = _jetzt()
        with self.lock:
            kurs = self.laden(kennung)
            aufgabe = next(a for a in kurs["aufgaben"] if a["id"] == aufgabe_id)
            aufgabe["gesehen"] = int(aufgabe.get("gesehen", 0)) + 1
            if richtig:
                aufgabe["richtig"] = int(aufgabe.get("richtig", 0)) + 1
                aufgabe["stufe"] = min(len(INTERVALLE) - 1, int(aufgabe.get("stufe", 0)) + 1)
            elif teilweise:
                aufgabe["stufe"] = max(0, int(aufgabe.get("stufe", 0)))
            else:
                aufgabe["falsch"] = int(aufgabe.get("falsch", 0)) + 1
                aufgabe["stufe"] = 0
            aufgabe["faellig"] = jetzt + INTERVALLE[aufgabe["stufe"]]
            kurs.setdefault("verlauf", []).append({"zeit": jetzt, "aufgabe": aufgabe_id, "richtig": richtig, "teilweise": teilweise})
            del kurs["verlauf"][:-500]
            self._beherrschung(kurs, aufgabe["lektion"])
            heute = date.today().isoformat()
            if kurs.get("letzter_tag") != heute:
                gestern = date.fromordinal(date.today().toordinal() - 1).isoformat()
                kurs["serie"] = int(kurs.get("serie", 0)) + 1 if kurs.get("letzter_tag") == gestern else 1
                kurs["letzter_tag"] = heute
            kurs["zuletzt"] = jetzt
            self.speichern(kurs)
            lektion = next((l for l in kurs["lektionen"] if l["nr"] == aufgabe["lektion"]), {})
        return {
            "richtig": richtig,
            "teilweise": teilweise,
            "feedback": feedback,
            "loesung": aufgabe["loesung"],
            "erklaerung": aufgabe.get("erklaerung", ""),
            "beherrschung": lektion.get("beherrschung", 0),
            "naechste_wiederholung": aufgabe["faellig"],
            "serie": kurs.get("serie", 0),
        }

    def _beherrschung(self, kurs: dict, lektion: int) -> None:
        aufgaben = [a for a in kurs.get("aufgaben", []) if a.get("lektion") == lektion]
        if not aufgaben:
            return
        punkte = 0.0
        for aufgabe in aufgaben:
            if not aufgabe.get("gesehen"):
                continue
            versuche = aufgabe.get("richtig", 0) + aufgabe.get("falsch", 0)
            quote = aufgabe.get("richtig", 0) / versuche if versuche else 0
            punkte += min(1.0, 0.55 * quote + 0.45 * min(1.0, aufgabe.get("stufe", 0) / 3))
        wert = round(punkte / max(len(aufgaben), 4), 2)
        for eintrag in kurs.get("lektionen", []):
            if eintrag["nr"] == lektion:
                eintrag["beherrschung"] = min(1.0, wert)

    @staticmethod
    def _auszug(material: str, frage: str, grenze: int = 14_000) -> str:
        woerter = {w for w in re.findall(r"\w{4,}", frage.casefold())}
        absaetze = [a for a in re.split(r"\n\s*\n", material) if a.strip()]
        if not woerter or not absaetze:
            return material[:grenze]
        bewertet = sorted(range(len(absaetze)), key=lambda i: -sum(1 for w in woerter if w in absaetze[i].casefold()))
        gewaehlt: list[int] = []
        laenge = 0
        for index in bewertet:
            if laenge > grenze:
                break
            gewaehlt.append(index)
            laenge += len(absaetze[index])
        return "\n\n".join(absaetze[i] for i in sorted(gewaehlt))[:grenze]

    def _stellen(self, kennung: str, frage: str, grenze: int = 3) -> str:
        material = self.material(kennung)
        if not material:
            return ""
        woerter = {w for w in re.findall(r"\w{4,}", frage.casefold())}
        if not woerter:
            return material[:3000]
        absaetze = [a for a in re.split(r"\n\s*\n", material) if a.strip()]
        bewertet = sorted(absaetze, key=lambda a: -sum(1 for w in woerter if w in a.casefold()))
        return "\n\n".join(a[:1800] for a in bewertet[:grenze] if any(w in a.casefold() for w in woerter))

    async def fragen(self, kennung: str, text: str, lektion: int | None = None) -> dict:
        kurs = self.laden(kennung)
        text = str(text or "").strip()
        if not text:
            raise ValueError("Was möchtest du wissen?")
        eintrag = next((l for l in kurs.get("lektionen", []) if l["nr"] == lektion), None) if lektion else None
        stellen = self._stellen(kennung, text)
        verlauf = "\n".join(f"{n['rolle']}: {n['text'][:700]}" for n in kurs.get("chat", [])[-10:])
        kontext = (
            f"KURS: {kurs.get('titel')} (Fach: {kurs.get('fach')}, Ziel: {kurs.get('ziel') or '-'}, Niveau: {kurs.get('niveau') or '-'})\n"
            f"ÜBERBLICK: {kurs.get('zusammenfassung', '')}\n"
            + (f"AKTUELLE LEKTION: {eintrag['titel']}\n{eintrag.get('inhalt', '')[:3500]}\n" if eintrag else "")
            + (f"PASSENDE STELLEN AUS DEM MATERIAL:\n{stellen}\n" if stellen else "")
            + (f"BISHERIGES GESPRÄCH:\n{verlauf}\n" if verlauf else "")
        )
        system = LEHRER_ROLLE + "\nBeantworte die Frage verständlich, mit einem Beispiel und einer kurzen Kontrollfrage am Ende, wenn das beim Lernen hilft. Höchstens 250 Wörter. Kein Markdown ausser **fett** und kurzen Listen."
        antwort = await self._modell(system, kontext + "\nFRAGE DES LERNENDEN: " + text, 1800, 0.5)
        antwort = antwort.strip() or "Da bin ich gerade nicht durchgekommen – frag mich bitte noch einmal."
        with self.lock:
            kurs = self.laden(kennung)
            kurs.setdefault("chat", []).extend([{"rolle": "du", "text": text[:2000], "zeit": _jetzt()}, {"rolle": "lehrer", "text": antwort[:6000], "zeit": _jetzt()}])
            del kurs["chat"][:-60]
            kurs["zuletzt"] = _jetzt()
            self.speichern(kurs)
        return {"antwort": antwort}

    async def erklaeren(self, kennung: str, lektion: int, art: str = "einfacher") -> dict:
        kurs = self.laden(kennung)
        eintrag = next((l for l in kurs.get("lektionen", []) if l["nr"] == lektion), None)
        if eintrag is None:
            raise ValueError("Diese Lektion gibt es nicht.")
        wunsch = {"einfacher": "Erkläre die Lektion noch einmal ganz einfach, mit Alltagsvergleich und Schritt-für-Schritt-Beispiel.", "beispiele": "Gib 5 neue, abwechslungsreiche Beispiele mit kurzer Begründung.", "merkzettel": "Schreibe einen kompakten Merkzettel (Spickzettel) mit den wichtigsten Regeln, Formeln oder Vokabeln, Eselsbrücken inklusive."}.get(art, art)
        antwort = await self._modell(LEHRER_ROLLE, f"Lektion „{eintrag['titel']}“:\n{eintrag.get('inhalt', '')}\n\n{wunsch}", 2500, 0.5)
        return {"text": antwort.strip()}

    def probetest(self, kennung: str, anzahl: int = 10) -> dict:
        kurs = self.laden(kennung)
        if kurs.get("status") != "bereit":
            raise ValueError("Der Kurs ist noch nicht fertig.")
        nach_lektion: dict[int, list[dict]] = {}
        for aufgabe in kurs.get("aufgaben", []):
            nach_lektion.setdefault(aufgabe["lektion"], []).append(aufgabe)
        gewicht = {l["nr"]: 1.6 - l.get("beherrschung", 0) for l in kurs.get("lektionen", [])}
        auswahl: list[dict] = []
        reihenfolge = sorted(nach_lektion, key=lambda nr: -gewicht.get(nr, 1))
        while len(auswahl) < anzahl and any(nach_lektion.values()):
            for nr in reihenfolge:
                if nach_lektion[nr] and len(auswahl) < anzahl:
                    random.shuffle(nach_lektion[nr])
                    nach_lektion[nr].sort(key=lambda a: -a.get("schwierigkeit", 1))
                    auswahl.append(nach_lektion[nr].pop(0))
        return {"aufgaben": [self._oeffentlich(a) for a in auswahl], "minuten": max(5, len(auswahl) * 2)}

    def entfernen(self, kennung: str) -> dict:
        import shutil

        ordner = self._pfad(kennung)
        if ordner.exists():
            shutil.rmtree(ordner, ignore_errors=True)
        return {"ok": True}

    def textstand(self) -> str:
        kurse = self.liste()
        if not kurse:
            return "Noch kein Kurs. Sag mir, was du lernst, und gib mir deine Lerndateien – dann baue ich dir einen Kurs."
        zeilen = []
        for kurs in kurse[:8]:
            zeilen.append(f"- {kurs['titel']} ({kurs['status']}): {round(kurs['beherrschung'] * 100)} % sicher, {kurs['faellig']} Wiederholungen fällig, {kurs['lektionen']} Lektionen, ID {kurs['id']}")
        return "\n".join(zeilen)


def _voller_text(pfad: Path, vorhanden: str) -> str:
    endung = pfad.suffix.lower()
    try:
        if endung == ".pdf":
            from pypdf import PdfReader

            leser = PdfReader(str(pfad))
            return "\n".join((seite.extract_text() or "") for seite in leser.pages[:400])[:MATERIAL_MAX]
        if endung == ".docx":
            import docx

            dokument = docx.Document(str(pfad))
            return "\n".join(a.text for a in dokument.paragraphs)[:MATERIAL_MAX]
        if endung == ".pptx":
            from pptx import Presentation

            folien = []
            for nummer, folie in enumerate(Presentation(str(pfad)).slides, 1):
                texte = [f.text_frame.text for f in folie.shapes if getattr(f, "has_text_frame", False) and f.has_text_frame]
                folien.append(f"Folie {nummer}: " + "\n".join(t for t in texte if t.strip()))
            return "\n\n".join(folien)[:MATERIAL_MAX]
        if endung in (".txt", ".md", ".csv", ".json", ".html", ".htm"):
            return pfad.read_text(encoding="utf-8", errors="replace")[:MATERIAL_MAX]
    except Exception as fehler:
        leise(fehler, "lehrer/volltext")
    return vorhanden


_dienst: LehrerService | None = None


def get_lehrer_service() -> LehrerService:
    global _dienst
    if _dienst is None:
        _dienst = LehrerService()
    return _dienst
