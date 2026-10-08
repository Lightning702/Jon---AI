from __future__ import annotations

import csv
import datetime as dt
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from app.core.fehler import leise

MARKDOWN = {"md", "markdown"}
HTML = {"html", "htm"}
BILD = {"png", "jpg", "jpeg", "webp", "gif", "bmp"}
CODE = {
    "txt", "log", "ini", "cfg", "conf", "toml", "yaml", "yml", "json", "xml", "py", "js", "ts",
    "tsx", "jsx", "css", "scss", "java", "kt", "kts", "c", "cpp", "h", "hpp", "cs", "go", "rs",
    "sh", "bat", "ps1", "sql", "env", "gradle", "properties", "vue", "php", "rb", "swift",
}
TEXT_MAX = 400_000
ZEILEN_MAX = 300
SPALTEN_MAX = 40
FOLIEN_MAX = 400
EMU_PRO_PUNKT = 12700


def _lesen(pfad: Path) -> tuple[str, bool]:
    with pfad.open("rb") as datei:
        roh = datei.read(TEXT_MAX + 1)
    return roh[:TEXT_MAX].decode("utf-8", errors="replace"), len(roh) > TEXT_MAX


def _zelle(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "WAHR" if wert else "FALSCH"
    if isinstance(wert, float):
        return str(int(wert)) if wert.is_integer() else f"{wert:.6g}"
    if isinstance(wert, (dt.datetime, dt.date, dt.time)):
        return wert.isoformat(sep=" ") if isinstance(wert, dt.datetime) else wert.isoformat()
    return str(wert)


def _stutzen(zeilen: list[list[str]]) -> list[list[str]]:
    while zeilen and not any(z.strip() for z in zeilen[-1]):
        zeilen.pop()
    breite = 0
    for zeile in zeilen:
        for index in range(len(zeile) - 1, -1, -1):
            if zeile[index].strip():
                breite = max(breite, index + 1)
                break
    return [zeile[:breite] + [""] * (breite - len(zeile[:breite])) for zeile in zeilen]


def _csv(pfad: Path) -> dict:
    text, gekuerzt = _lesen(pfad)
    try:
        dialekt = csv.Sniffer().sniff(text[:4000], delimiters=";,\t|")
    except csv.Error:
        dialekt = csv.excel
    zeilen = []
    for nummer, zeile in enumerate(csv.reader(io.StringIO(text), dialekt)):
        if nummer >= ZEILEN_MAX:
            gekuerzt = True
            break
        zeilen.append([str(z) for z in zeile[:SPALTEN_MAX]])
    return {"name": pfad.stem, "zeilen": _stutzen(zeilen), "gekuerzt": gekuerzt}


def _xlsx(pfad: Path) -> list[dict]:
    from openpyxl import load_workbook

    mappe = load_workbook(str(pfad), read_only=True, data_only=True)
    try:
        blaetter = []
        for blatt in mappe.worksheets[:12]:
            zeilen = []
            gekuerzt = False
            for nummer, zeile in enumerate(blatt.iter_rows(values_only=True)):
                if nummer >= ZEILEN_MAX:
                    gekuerzt = True
                    break
                zeilen.append([_zelle(w) for w in list(zeile)[:SPALTEN_MAX]])
            blaetter.append({"name": blatt.title, "zeilen": _stutzen(zeilen), "gekuerzt": gekuerzt})
        return blaetter
    finally:
        mappe.close()


def _laufe(absatz) -> str:
    teile = []
    for lauf in absatz.runs:
        text = lauf.text
        if not text:
            continue
        if lauf.bold and text.strip():
            text = f"**{text.strip()}** " if text.endswith(" ") else f"**{text.strip()}**"
        elif lauf.italic and text.strip():
            text = f"*{text.strip()}*"
        teile.append(text)
    return "".join(teile) or absatz.text


def _tabelle_markdown(zeilen: list[list[str]]) -> str:
    if not zeilen:
        return ""
    sauber = [[z.replace("|", "\\|").replace("\n", " ").strip() for z in zeile] for zeile in zeilen]
    breite = max(len(z) for z in sauber)
    sauber = [z + [""] * (breite - len(z)) for z in sauber]
    kopf = "| " + " | ".join(sauber[0]) + " |"
    trenner = "| " + " | ".join("---" for _ in range(breite)) + " |"
    rumpf = ["| " + " | ".join(z) + " |" for z in sauber[1:]]
    return "\n".join([kopf, trenner, *rumpf])


def _docx(pfad: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    dokument = Document(str(pfad))
    zeilen: list[str] = []
    for kind in dokument.element.body.iterchildren():
        name = kind.tag.rsplit("}", 1)[-1]
        if name == "p":
            absatz = Paragraph(kind, dokument)
            text = _laufe(absatz).strip()
            if not text:
                continue
            stil = (absatz.style.name if absatz.style is not None else "").lower()
            if stil.startswith("title") or stil.startswith("titel"):
                zeilen.append(f"# {absatz.text.strip()}")
            elif "heading" in stil or "überschrift" in stil or "uberschrift" in stil:
                ziffer = next((int(z) for z in stil if z.isdigit()), 2)
                zeilen.append(f"{'#' * min(6, ziffer + 1)} {absatz.text.strip()}")
            elif "list number" in stil or "nummer" in stil:
                zeilen.append(f"1. {text}")
            elif "list" in stil or "aufzählung" in stil or "liste" in stil:
                zeilen.append(f"- {text}")
            else:
                zeilen.append(text)
        elif name == "tbl":
            tabelle = Table(kind, dokument)
            daten = [[zelle.text for zelle in reihe.cells] for reihe in tabelle.rows][:ZEILEN_MAX]
            zeilen.append(_tabelle_markdown(daten))
        if sum(len(z) for z in zeilen) > TEXT_MAX:
            break
    return "\n\n".join(zeilen)


def docx_elemente(dokument) -> list:
    return [k for k in dokument.element.body.iterchildren() if k.tag.rsplit("}", 1)[-1] in ("p", "tbl")]


def _stilart(absatz) -> str:
    stil = (absatz.style.name if absatz.style is not None else "").lower()
    if stil.startswith("title") or stil.startswith("titel"):
        return "titel"
    if "heading" in stil or "überschrift" in stil or "uberschrift" in stil:
        ziffer = next((int(z) for z in stil if z.isdigit()), 1)
        return f"h{min(6, max(1, ziffer))}"
    if "list number" in stil or "nummer" in stil:
        return "nummer"
    if "list" in stil or "aufzählung" in stil or "liste" in stil:
        return "liste"
    if "quote" in stil or "zitat" in stil:
        return "zitat"
    return "absatz"


def docx_bloecke(pfad: Path) -> list[dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    dokument = Document(str(pfad))
    bloecke: list[dict] = []
    laenge = 0
    for nummer, kind in enumerate(docx_elemente(dokument)):
        if kind.tag.endswith("}tbl"):
            tabelle = Table(kind, dokument)
            zeilen = [[zelle.text for zelle in reihe.cells] for reihe in tabelle.rows][:ZEILEN_MAX]
            bloecke.append({"nr": nummer, "art": "tabelle", "zeilen": zeilen})
            continue
        absatz = Paragraph(kind, dokument)
        bilder = [b for b in kind.xpath(".//a:blip/@r:embed")]
        eintrag = {"nr": nummer, "art": _stilart(absatz), "text": _laufe(absatz).strip()}
        if absatz.alignment is not None:
            eintrag["ausrichtung"] = "center" if "CENTER" in str(absatz.alignment) else "right" if "RIGHT" in str(absatz.alignment) else "left"
        if bilder:
            eintrag["bilder"] = bilder[:6]
        bloecke.append(eintrag)
        laenge += len(eintrag["text"])
        if laenge > TEXT_MAX or len(bloecke) > 4000:
            break
    return bloecke


def docxbild(pfad: Path, kennung: str) -> tuple[bytes, str]:
    from docx import Document

    teil = Document(str(pfad)).part.related_parts[kennung]
    return teil.blob, getattr(teil, "content_type", "image/png")


def _odt(pfad: Path) -> str:
    from odf import teletype, text
    from odf.opendocument import load

    dokument = load(str(pfad))
    zeilen = []
    for element in dokument.getElementsByType(text.H) + dokument.getElementsByType(text.P):
        inhalt = teletype.extractText(element).strip()
        if inhalt:
            zeilen.append(f"## {inhalt}" if element.qname[1] == "h" else inhalt)
    return "\n\n".join(zeilen)[:TEXT_MAX]


def _farbe(farbe) -> str | None:
    try:
        if farbe is not None and farbe.type is not None and farbe.rgb is not None:
            return f"#{farbe.rgb}"
    except Exception:
        return None
    return None


def _fuellung(form) -> str | None:
    try:
        from pptx.enum.dml import MSO_FILL

        if form.fill.type == MSO_FILL.SOLID:
            return _farbe(form.fill.fore_color)
    except Exception:
        return None
    return None


def _hintergrund(folie) -> str | None:
    for quelle in (folie, getattr(folie, "slide_layout", None), getattr(getattr(folie, "slide_layout", None), "slide_master", None)):
        if quelle is None:
            continue
        try:
            wert = _fuellung(quelle.background)
            if wert:
                return wert
        except Exception:
            continue
    return None


def _absaetze(rahmen) -> list[dict]:
    ergebnis = []
    for absatz in rahmen.paragraphs:
        text = "".join(lauf.text for lauf in absatz.runs) or absatz.text
        if not text.strip():
            ergebnis.append({"text": "", "groesse": 10})
            continue
        lauf = next((l for l in absatz.runs if l.text.strip()), None)
        schrift = lauf.font if lauf is not None else absatz.font
        groesse = schrift.size.pt if schrift is not None and schrift.size is not None else (absatz.font.size.pt if absatz.font.size is not None else None)
        ausrichtung = str(absatz.alignment).split(".")[-1].split(" ")[0].lower() if absatz.alignment is not None else ""
        ergebnis.append({
            "text": text,
            "groesse": groesse,
            "fett": bool(schrift.bold) if schrift is not None and schrift.bold is not None else None,
            "farbe": _farbe(schrift.color) if schrift is not None else None,
            "ausrichtung": "center" if "center" in ausrichtung else "right" if "right" in ausrichtung else "left",
            "ebene": absatz.level,
        })
    return ergebnis


def _pptx(pfad: Path) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE

    praesentation = Presentation(str(pfad))
    breite = praesentation.slide_width or 12192000
    hoehe = praesentation.slide_height or 6858000
    folien = []
    for nummer, folie in enumerate(praesentation.slides):
        if nummer >= FOLIEN_MAX:
            break
        formen = []
        for index, form in enumerate(folie.shapes):
            if form.left is None or form.top is None or form.width is None or form.height is None:
                continue
            eintrag = {
                "index": index,
                "x": round(form.left / breite * 100, 3),
                "y": round(form.top / hoehe * 100, 3),
                "b": round(form.width / breite * 100, 3),
                "h": round(form.height / hoehe * 100, 3),
                "drehung": round(float(getattr(form, "rotation", 0) or 0), 2),
                "name": str(getattr(form, "name", ""))[:80],
            }
            try:
                if form.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    eintrag["art"] = "bild"
                elif getattr(form, "has_table", False) and form.has_table:
                    eintrag["art"] = "tabelle"
                    eintrag["zeilen"] = [[zelle.text for zelle in reihe.cells] for reihe in form.table.rows][:40]
                elif getattr(form, "has_chart", False) and form.has_chart:
                    eintrag["art"] = "diagramm"
                    diagramm = form.chart
                    eintrag["titel"] = diagramm.chart_title.text_frame.text if diagramm.has_title else ""
                    reihen = []
                    for plot in diagramm.plots:
                        kategorien = [str(k) for k in plot.categories]
                        for reihe in plot.series:
                            reihen.append({"name": reihe.name, "werte": [float(w or 0) for w in reihe.values], "kategorien": kategorien})
                    eintrag["reihen"] = reihen[:6]
                else:
                    eintrag["art"] = "form"
                    eintrag["fuellung"] = _fuellung(form)
                    try:
                        typ = form.auto_shape_type
                        eintrag["rund"] = typ in (MSO_AUTO_SHAPE_TYPE.OVAL,)
                        eintrag["ecken"] = typ in (MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,)
                    except Exception:
                        pass
                    if form.has_text_frame and form.text_frame.text.strip():
                        eintrag["absaetze"] = _absaetze(form.text_frame)
            except Exception as fehler:
                leise(fehler, "services/dateiansicht_service")
                continue
            formen.append(eintrag)
        notizen = ""
        try:
            if folie.has_notes_slide:
                notizen = folie.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notizen = ""
        folien.append({"nummer": nummer + 1, "hintergrund": _hintergrund(folie), "formen": formen, "notizen": notizen[:2000]})
    return {"verhaeltnis": round(breite / hoehe, 4), "hoehe_pt": round(hoehe / EMU_PRO_PUNKT, 2), "folien": folien, "gesamt": len(praesentation.slides)}


def folienbild(pfad: Path, folie: int, form: int) -> tuple[bytes, str]:
    from PIL import Image
    from pptx import Presentation

    praesentation = Presentation(str(pfad))
    ziel = praesentation.slides[folie - 1].shapes[form]
    roh = ziel.image.blob
    with Image.open(io.BytesIO(roh)) as bild:
        bild.thumbnail((1200, 1200))
        ausgabe = io.BytesIO()
        if bild.mode in ("RGBA", "LA", "P"):
            bild.convert("RGBA").save(ausgabe, format="PNG", optimize=True)
            return ausgabe.getvalue(), "image/png"
        bild.convert("RGB").save(ausgabe, format="JPEG", quality=84)
        return ausgabe.getvalue(), "image/jpeg"


def _pdf(pfad: Path) -> dict:
    from pypdf import PdfReader

    leser = PdfReader(str(pfad))
    seiten = len(leser.pages)
    if shutil.which("pdftoppm"):
        return {"seiten": seiten, "bilder": True}
    texte = []
    for seite in leser.pages[:40]:
        try:
            texte.append(seite.extract_text() or "")
        except Exception:
            texte.append("")
    return {"seiten": seiten, "bilder": False, "texte": texte}


def pdfseite(pfad: Path, seite: int) -> bytes:
    programm = shutil.which("pdftoppm")
    if not programm:
        raise RuntimeError("Seitenbilder sind auf diesem Gerät nicht verfügbar.")
    with tempfile.TemporaryDirectory() as ordner:
        ziel = Path(ordner) / "seite"
        subprocess.run(
            [programm, "-jpeg", "-jpegopt", "quality=82", "-scale-to", "1200", "-f", str(seite), "-l", str(seite), str(pfad), str(ziel)],
            check=True, capture_output=True, timeout=60,
        )
        bilder = sorted(Path(ordner).glob("seite*.jpg"))
        if not bilder:
            raise RuntimeError("Seite konnte nicht gerendert werden.")
        return bilder[0].read_bytes()


def ansehen(pfad: Path) -> dict:
    endung = pfad.suffix.lower().lstrip(".")
    info = pfad.stat()
    basis = {"name": pfad.name, "endung": endung, "groesse": info.st_size, "pfad": str(pfad), "geaendert": info.st_mtime}
    try:
        if endung in MARKDOWN:
            text, gekuerzt = _lesen(pfad)
            return {**basis, "art": "markdown", "text": text, "gekuerzt": gekuerzt}
        if endung in HTML:
            text, gekuerzt = _lesen(pfad)
            return {**basis, "art": "html", "text": text, "gekuerzt": gekuerzt}
        if endung == "csv" or endung == "tsv":
            return {**basis, "art": "tabelle", "blaetter": [_csv(pfad)]}
        if endung in CODE or not endung:
            text, gekuerzt = _lesen(pfad)
            return {**basis, "art": "code", "sprache": endung or "text", "text": text, "gekuerzt": gekuerzt}
        if endung == "docx":
            return {**basis, "art": "markdown", "text": _docx(pfad), "bloecke": docx_bloecke(pfad)}
        if endung == "odt":
            return {**basis, "art": "markdown", "text": _odt(pfad)}
        if endung in ("xlsx", "xlsm"):
            return {**basis, "art": "tabelle", "blaetter": _xlsx(pfad)}
        if endung == "pptx":
            return {**basis, "art": "folien", **_pptx(pfad)}
        if endung == "pdf":
            return {**basis, "art": "pdf", **_pdf(pfad)}
        if endung in BILD:
            return {**basis, "art": "bild"}
    except Exception as fehler:
        leise(fehler, "services/dateiansicht_service")
        return {**basis, "art": "fehler", "fehler": f"Die Datei ließ sich nicht öffnen: {fehler}"}
    return {**basis, "art": "unbekannt"}
