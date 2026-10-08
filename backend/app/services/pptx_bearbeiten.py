from __future__ import annotations

import copy
from pathlib import Path
from typing import Any

from app.core.fehler import leise

POSITIONEN = {
    "rechts": (0.52, 0.2, 0.43, 0.68),
    "links": (0.05, 0.2, 0.43, 0.68),
    "vollbild": (0.0, 0.0, 1.0, 1.0),
    "hintergrund": (0.0, 0.0, 1.0, 1.0),
    "mitte": (0.2, 0.2, 0.6, 0.65),
    "oben": (0.05, 0.05, 0.9, 0.45),
    "unten": (0.05, 0.5, 0.9, 0.45),
    "klein": (0.7, 0.68, 0.25, 0.25),
}
BILDFELDER = ("image", "bild", "image_url", "bild_url", "bild_suche", "image_search", "bild_prompt", "image_prompt", "pfad")


def _zahl(wert, standard: float) -> float:
    try:
        zahl = float(wert)
    except (TypeError, ValueError):
        return standard
    return zahl / 100 if zahl > 1.0 else zahl


def _rechteck(aktion: dict, breite: int, hoehe: int, standard: str = "rechts") -> tuple[int, int, int, int]:
    position = aktion.get("position")
    if isinstance(position, dict):
        werte = position
    else:
        werte = {}
        x, y, b, h = POSITIONEN.get(str(position or standard).strip().lower(), POSITIONEN[standard])
        werte = {"x": x, "y": y, "b": b, "h": h}
    for schluessel in ("x", "y", "b", "h"):
        if schluessel in aktion:
            werte[schluessel] = aktion[schluessel]
    x = _zahl(werte.get("x"), 0.1)
    y = _zahl(werte.get("y"), 0.1)
    b = _zahl(werte.get("b", werte.get("breite")), 0.4)
    h = _zahl(werte.get("h", werte.get("hoehe")), 0.4)
    return int(x * breite), int(y * hoehe), int(b * breite), int(h * hoehe)


def text_setzen(rahmen, text: str) -> None:
    zeilen = str(text).split("\n")
    absaetze = rahmen.paragraphs
    erste = absaetze[0]
    for alt in absaetze[1:]:
        alt._p.getparent().remove(alt._p)
    vorlage = copy.deepcopy(erste._p)
    _absatz_setzen(erste, zeilen[0])
    letzter = erste._p
    for zeile in zeilen[1:]:
        neu = copy.deepcopy(vorlage)
        letzter.addnext(neu)
        letzter = neu
        from pptx.text.text import _Paragraph

        _absatz_setzen(_Paragraph(neu, erste._parent), zeile)


def _absatz_setzen(absatz, text: str) -> None:
    laeufe = absatz.runs
    if not laeufe:
        absatz.text = text
        return
    laeufe[0].text = text
    for rest in laeufe[1:]:
        rest._r.getparent().remove(rest._r)


def _ersetzen_in_rahmen(rahmen, suchen: str, ersetzen: str) -> int:
    anzahl = 0
    for absatz in rahmen.paragraphs:
        for lauf in absatz.runs:
            if suchen in lauf.text:
                anzahl += lauf.text.count(suchen)
                lauf.text = lauf.text.replace(suchen, ersetzen)
        gesamt = "".join(l.text for l in absatz.runs)
        if suchen in gesamt and absatz.runs:
            anzahl += gesamt.count(suchen)
            _absatz_setzen(absatz, gesamt.replace(suchen, ersetzen))
    return anzahl


def _rahmen(folie):
    for form in folie.shapes:
        if getattr(form, "has_text_frame", False) and form.has_text_frame:
            yield form.text_frame
        if getattr(form, "has_table", False) and form.has_table:
            for reihe in form.table.rows:
                for zelle in reihe.cells:
                    yield zelle.text_frame


def _folie_kopieren(praesentation, quelle) -> Any:
    layout = quelle.slide_layout
    neu = praesentation.slides.add_slide(layout)
    for form in list(neu.shapes):
        form._element.getparent().remove(form._element)
    zuordnung: dict[str, str] = {}
    for kennung, beziehung in quelle.part.rels.items():
        if "notesSlide" in beziehung.reltype or "slideLayout" in beziehung.reltype:
            continue
        try:
            if beziehung.is_external:
                zuordnung[kennung] = neu.part.rels.get_or_add_ext_rel(beziehung.reltype, beziehung.target_ref)
            else:
                zuordnung[kennung] = neu.part.rels.get_or_add(beziehung.reltype, beziehung.target_part)
        except Exception as fehler:
            leise(fehler, "services/pptx_bearbeiten")
    baum = neu.shapes._spTree
    for form in quelle.shapes:
        element = copy.deepcopy(form._element)
        for knoten in element.iter():
            for attribut, wert in list(knoten.attrib.items()):
                if attribut.endswith("}embed") or attribut.endswith("}id") or attribut.endswith("}link"):
                    if wert in zuordnung:
                        knoten.set(attribut, zuordnung[wert])
        baum.append(element)
    hintergrund = quelle._element.cSld.bg
    if hintergrund is not None:
        neu._element.cSld.insert(0, copy.deepcopy(hintergrund))
    return neu


def _einordnen(praesentation, folie, nach: int) -> None:
    liste = praesentation.slides._sldIdLst
    eintraege = list(liste)
    eigener = next(e for e in eintraege if praesentation.part.related_part(e.rId) is folie.part)
    liste.remove(eigener)
    ziel = max(0, min(int(nach), len(liste)))
    liste.insert(ziel, eigener)


def _leere_folie(praesentation):
    layouts = list(praesentation.slide_layouts)
    leer = next((l for l in layouts if len(l.placeholders) == 0), None)
    folie = praesentation.slides.add_slide(leer or layouts[min(6, len(layouts) - 1)])
    for form in list(folie.placeholders):
        form._element.getparent().remove(form._element)
    return folie


async def bearbeiten(pfad: str, aktionen: list[dict], theme: str = "") -> dict:
    import asyncio

    from app.services.bildquelle import besorgen

    quelle = Path(str(pfad or "")).expanduser()
    if not quelle.is_file() or quelle.suffix.lower() != ".pptx":
        return {"error": f"PowerPoint nicht gefunden: {quelle}"}
    if not isinstance(aktionen, list) or not aktionen:
        return {"error": "aktionen fehlt. Gib eine Liste von Bearbeitungsschritten an."}
    for aktion in aktionen:
        if isinstance(aktion, dict) and any(str(aktion.get(f) or "").strip() for f in BILDFELDER if f != "pfad"):
            art = str(aktion.get("aktion", "")).lower()
            if art == "folie_neu":
                from app.services.pptx_normalisieren import folie as folie_normalisieren

                aktion.update({k: v for k, v in folie_normalisieren(dict(aktion.get("inhalt") or aktion)).items() if k in BILDFELDER})
            if art in ("bild", "bild_einfuegen", "bild_ersetzen", "folie_neu"):
                breit = str(aktion.get("layout", "")).lower() not in ("image", "bild")
                aktion["_bild"] = await besorgen(aktion, breit)
    return await asyncio.to_thread(_anwenden, quelle, aktionen, theme)


def _anwenden(quelle: Path, aktionen: list[dict], theme: str) -> dict:
    from pptx import Presentation
    from pptx.util import Pt

    from app.services import pptx_service as gestaltung

    praesentation = Presentation(str(quelle))
    breite = praesentation.slide_width
    hoehe = praesentation.slide_height
    folien = list(praesentation.slides)
    formen = {nummer: list(folie.shapes) for nummer, folie in enumerate(folien, start=1)}
    farben = gestaltung._theme(theme or gestaltung.DEFAULT_THEME)
    erledigt: list[str] = []
    fehler: list[str] = []

    def folie_von(aktion: dict):
        nummer = int(aktion.get("folie") or 0)
        if nummer < 1 or nummer > len(folien):
            raise ValueError(f"Folie {nummer} gibt es nicht (1 bis {len(folien)}).")
        return nummer, folien[nummer - 1]

    def form_von(aktion: dict):
        nummer, folie = folie_von(aktion)
        index = aktion.get("element", aktion.get("form"))
        if index is None:
            raise ValueError("element fehlt (Nummer des Elements auf der Folie, siehe read_pptx).")
        liste = formen[nummer]
        if int(index) < 0 or int(index) >= len(liste):
            raise ValueError(f"Element {index} gibt es auf Folie {nummer} nicht (0 bis {len(liste) - 1}).")
        return nummer, folie, liste[int(index)]

    for schritt, aktion in enumerate(aktionen, start=1):
        if not isinstance(aktion, dict):
            fehler.append(f"Schritt {schritt}: kein Objekt")
            continue
        art = str(aktion.get("aktion", "")).strip().lower()
        try:
            if art == "text":
                nummer, folie, form = form_von(aktion)
                if getattr(form, "has_table", False) and form.has_table and ("zeile" in aktion or "spalte" in aktion):
                    zelle = form.table.cell(int(aktion.get("zeile", 0)), int(aktion.get("spalte", 0)))
                    text_setzen(zelle.text_frame, str(aktion.get("text", "")))
                elif not form.has_text_frame:
                    raise ValueError("Dieses Element hat keinen Text.")
                else:
                    text_setzen(form.text_frame, str(aktion.get("text", "")))
                erledigt.append(f"Folie {nummer}: Text von Element {aktion.get('element', aktion.get('form'))} geändert")
            elif art == "ersetzen":
                suchen, ersetzen = str(aktion.get("suchen", "")), str(aktion.get("ersetzen", ""))
                if not suchen:
                    raise ValueError("suchen fehlt.")
                ziele = [folie_von(aktion)[1]] if aktion.get("folie") else list(praesentation.slides)
                anzahl = sum(_ersetzen_in_rahmen(rahmen, suchen, ersetzen) for folie in ziele for rahmen in _rahmen(folie))
                erledigt.append(f"„{suchen}“ {anzahl}× ersetzt")
            elif art in ("schrift", "format"):
                nummer, folie, form = form_von(aktion)
                for absatz in form.text_frame.paragraphs:
                    for lauf in absatz.runs:
                        if aktion.get("groesse"):
                            lauf.font.size = Pt(float(aktion["groesse"]))
                        if "fett" in aktion:
                            lauf.font.bold = bool(aktion["fett"])
                        if "kursiv" in aktion:
                            lauf.font.italic = bool(aktion["kursiv"])
                        if aktion.get("farbe"):
                            lauf.font.color.rgb = gestaltung._rgb(str(aktion["farbe"]))
                erledigt.append(f"Folie {nummer}: Schrift von Element {aktion.get('element')} angepasst")
            elif art == "notizen":
                nummer, folie = folie_von(aktion)
                folie.notes_slide.notes_text_frame.text = str(aktion.get("text", ""))
                erledigt.append(f"Folie {nummer}: Sprechernotizen gesetzt")
            elif art in ("bild", "bild_einfuegen", "bild_ersetzen"):
                nummer, folie = folie_von(aktion)
                bild = str(aktion.get("_bild") or aktion.get("pfad") or "")
                if not bild or not Path(bild).expanduser().is_file():
                    raise ValueError("Kein Bild gefunden. Gib image (Pfad oder URL), bild_suche oder bild_prompt an.")
                alt = None
                if aktion.get("element") is not None:
                    _, _, alt = form_von(aktion)
                if alt is not None:
                    x, y, b, h = alt.left, alt.top, alt.width, alt.height
                else:
                    x, y, b, h = _rechteck(aktion, breite, hoehe)
                emu = 914400
                ganz = str(aktion.get("einpassen", "")).lower() in ("ganz", "true", "1", "enthalten")
                form = gestaltung.einpassen(folie, bild, x / emu, y / emu, b / emu, h / emu, ganz)
                if form is None:
                    raise ValueError("Das Bild ließ sich nicht einfügen.")
                if alt is not None:
                    alt._element.addprevious(form._element)
                    alt._element.getparent().remove(alt._element)
                elif str(aktion.get("position", "")).lower() == "hintergrund":
                    folie.shapes._spTree.remove(form._element)
                    folie.shapes._spTree.insert(2, form._element)
                erledigt.append(f"Folie {nummer}: Bild {'ersetzt' if alt is not None else 'eingefügt'}")
            elif art == "textfeld":
                nummer, folie = folie_von(aktion)
                x, y, b, h = _rechteck(aktion, breite, hoehe, "unten")
                feld = folie.shapes.add_textbox(x, y, b, h)
                feld.text_frame.word_wrap = True
                text_setzen(feld.text_frame, str(aktion.get("text", "")))
                for absatz in feld.text_frame.paragraphs:
                    for lauf in absatz.runs:
                        lauf.font.size = Pt(float(aktion.get("groesse") or 20))
                        lauf.font.bold = bool(aktion.get("fett"))
                        lauf.font.color.rgb = gestaltung._rgb(str(aktion.get("farbe") or farben["dark"]))
                erledigt.append(f"Folie {nummer}: Textfeld hinzugefügt")
            elif art in ("element_loeschen", "loeschen"):
                nummer, folie, form = form_von(aktion)
                form._element.getparent().remove(form._element)
                erledigt.append(f"Folie {nummer}: Element {aktion.get('element')} entfernt")
            elif art in ("element_verschieben", "verschieben", "groesse"):
                nummer, folie, form = form_von(aktion)
                x, y, b, h = _rechteck({**{"x": form.left / breite, "y": form.top / hoehe, "b": form.width / breite, "h": form.height / hoehe}, **aktion, "position": {}}, breite, hoehe)
                form.left, form.top, form.width, form.height = x, y, b, h
                erledigt.append(f"Folie {nummer}: Element {aktion.get('element')} verschoben")
            elif art == "diagramm":
                nummer, folie, form = form_von(aktion)
                if not form.has_chart:
                    raise ValueError("Dieses Element ist kein Diagramm.")
                from pptx.chart.data import CategoryChartData

                daten = CategoryChartData()
                daten.categories = [str(k) for k in aktion.get("kategorien") or []]
                for reihe in aktion.get("reihen") or []:
                    daten.add_series(str(reihe.get("name", "")), [float(w or 0) for w in reihe.get("werte") or []])
                form.chart.replace_data(daten)
                erledigt.append(f"Folie {nummer}: Diagrammdaten ersetzt")
            elif art == "hintergrund":
                nummer, folie = folie_von(aktion)
                fuellung = folie.background.fill
                fuellung.solid()
                fuellung.fore_color.rgb = gestaltung._rgb(str(aktion.get("farbe", "FFFFFF")))
                erledigt.append(f"Folie {nummer}: Hintergrund gesetzt")
            elif art == "uebergang":
                from app.services.pptx_effekte import uebergang

                nummer, folie = folie_von(aktion)
                uebergang(folie, str(aktion.get("art") or aktion.get("uebergang") or "fade"), str(aktion.get("tempo", "mittel")))
                erledigt.append(f"Folie {nummer}: Übergang gesetzt")
            elif art == "folie_neu":
                from app.services.pptx_normalisieren import folie as folie_normalisieren

                daten = folie_normalisieren(dict(aktion.get("inhalt") or aktion))
                if aktion.get("_bild"):
                    daten["image"] = aktion["_bild"]
                layout = str(daten.get("layout", "bullets")).strip().lower()
                folie = _leere_folie(praesentation)
                gestaltung.LAYOUTS.get(layout, gestaltung._bullets_slide)(praesentation, folie, daten, farben)
                if daten.get("notes"):
                    folie.notes_slide.notes_text_frame.text = str(daten["notes"])
                nach = aktion.get("nach")
                if nach is not None:
                    _einordnen(praesentation, folie, int(nach))
                erledigt.append(f"Neue Folie ({layout}) {'nach Folie ' + str(nach) if nach is not None else 'am Ende'} eingefügt")
            elif art == "folie_loeschen":
                nummer, folie = folie_von(aktion)
                liste = praesentation.slides._sldIdLst
                eintrag = next(e for e in list(liste) if praesentation.part.related_part(e.rId) is folie.part)
                praesentation.part.drop_rel(eintrag.rId)
                liste.remove(eintrag)
                erledigt.append(f"Folie {nummer} gelöscht")
            elif art == "folie_verschieben":
                nummer, folie = folie_von(aktion)
                ziel = int(aktion.get("nach", aktion.get("position", nummer)))
                _einordnen(praesentation, folie, ziel)
                erledigt.append(f"Folie {nummer} verschoben")
            elif art == "folie_duplizieren":
                nummer, folie = folie_von(aktion)
                kopie = _folie_kopieren(praesentation, folie)
                position = list(praesentation.slides).index(folie) + 1
                _einordnen(praesentation, kopie, int(aktion.get("nach", position)))
                erledigt.append(f"Folie {nummer} dupliziert")
            else:
                raise ValueError(f"Unbekannte aktion '{art}'.")
        except Exception as exc:
            fehler.append(f"Schritt {schritt} ({art or '?'}): {exc}")
    if not erledigt:
        return {"error": "Nichts geändert. " + " | ".join(fehler)}
    ziel = quelle
    try:
        praesentation.save(str(quelle))
    except PermissionError:
        ziel = quelle.with_name(f"{quelle.stem} (bearbeitet).pptx")
        praesentation.save(str(ziel))
    ergebnis: dict = {"ok": True, "path": str(ziel), "folien": len(praesentation.slides), "erledigt": erledigt}
    if ziel != quelle:
        ergebnis["hinweis"] = "Die Originaldatei ist gerade in PowerPoint geöffnet und gesperrt - die Änderungen liegen in einer Kopie."
    if fehler:
        ergebnis["probleme"] = fehler
    return ergebnis


def _bild_sichern(quelle: Path, folie: int, index: int, form) -> str:
    try:
        from app.services.dateiraum_service import get_dateiraum_service

        bild = form.image
        ordner = get_dateiraum_service().ordner("Images") / "Vorlagen" / quelle.stem[:60]
        ordner.mkdir(parents=True, exist_ok=True)
        ziel = ordner / f"folie{folie}-element{index}.{bild.ext or 'png'}"
        if not ziel.exists():
            ziel.write_bytes(bild.blob)
        return str(ziel)
    except Exception as fehler:
        leise(fehler, "services/pptx_bearbeiten")
        return ""


def lesen(pfad: str, max_folien: int = 200) -> dict:
    from pptx import Presentation

    from app.services.kommentar_service import beschreibung, get_kommentar_service

    quelle = Path(str(pfad or "")).expanduser()
    if not quelle.is_file():
        return {"error": f"Datei nicht gefunden: {quelle}"}
    praesentation = Presentation(str(quelle))
    breite, hoehe = praesentation.slide_width, praesentation.slide_height
    folien = []
    for nummer, folie in enumerate(praesentation.slides, start=1):
        if nummer > max_folien:
            break
        elemente = []
        for index, form in enumerate(folie.shapes):
            eintrag: dict = {"element": index, "name": form.name}
            if form.left is not None and form.width is not None and breite:
                eintrag["lage"] = f"x {round(form.left / breite * 100)}% y {round(form.top / hoehe * 100)}% b {round(form.width / breite * 100)}% h {round(form.height / hoehe * 100)}%"
            if getattr(form, "has_text_frame", False) and form.has_text_frame and form.text_frame.text.strip():
                eintrag["text"] = form.text_frame.text.strip()[:1200]
            elif getattr(form, "has_table", False) and form.has_table:
                eintrag["tabelle"] = [[z.text for z in r.cells] for r in form.table.rows][:30]
            elif getattr(form, "has_chart", False) and form.has_chart:
                eintrag["diagramm"] = True
            elif form.shape_type == 13:
                eintrag["bild"] = True
                pfad = _bild_sichern(quelle, nummer, index, form)
                if pfad:
                    eintrag["bild_pfad"] = pfad
            else:
                continue
            elemente.append(eintrag)
        notizen = folie.notes_slide.notes_text_frame.text.strip() if folie.has_notes_slide else ""
        folien.append({"folie": nummer, "elemente": elemente, **({"notizen": notizen} if notizen else {})})
    kommentare = [
        {"id": k["id"], "wo": beschreibung(k.get("stelle") or {}), "stelle": k.get("stelle"), "text": k["text"]}
        for k in get_kommentar_service().liste(quelle, offen=True)
    ]
    ergebnis = {"path": str(quelle), "anzahl_folien": len(praesentation.slides), "folien": folien}
    if kommentare:
        ergebnis["offene_kommentare"] = kommentare
        ergebnis["hinweis"] = "Der Nutzer hat Kommentare hinterlassen. Setze sie mit edit_pptx um und markiere sie danach mit datei_kommentare (aktion erledigt) als erledigt."
    return ergebnis
