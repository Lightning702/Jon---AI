from __future__ import annotations

from app.services import pptx_service as basis

FARBEN = {
    "primary": "156082",
    "secondary": "0F9ED5",
    "accent": "E97132",
    "dark": "262626",
    "light": "FFFFFF",
    "muted": "595959",
    "soft": "F2F2F2",
    "tief": "0E2841",
    "stil": "hell",
}
TITEL = "Aptos Display"
TEXT = "Aptos"
TEXT_GROSS = 24
TEXT_KLEIN = 14
BREITE_FAKTOR = 0.49
ZEILEN_FAKTOR = 1.3


def _passt(zeilen: list[str], breite: float, groesse: int, hoehe: float) -> bool:
    pro_zeile = max(8, int(breite * 72 / (groesse * BREITE_FAKTOR)))
    gesamt = 0.0
    for zeile in zeilen:
        anzahl = max(1, -(-len(str(zeile)) // pro_zeile))
        gesamt += anzahl * groesse * ZEILEN_FAKTOR / 72 + 9 / 72
    return gesamt <= hoehe


def _wort_passt(text: str, breite: float, groesse: int) -> bool:
    laengstes = max((len(w) for w in str(text).split()), default=0)
    return laengstes * groesse * 0.55 / 72 <= breite


def aktiv(farben: dict) -> bool:
    return farben.get("stil") == "hell"


def _zoll(wert: float):
    from pptx.util import Inches

    return Inches(wert)


def _linie(slide, links: float, oben: float, breite: float, farbe: str, dicke: float = 2.75):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Pt

    linie = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _zoll(links), _zoll(oben), _zoll(links + breite), _zoll(oben))
    linie.line.color.rgb = basis._rgb(farbe)
    linie.line.width = Pt(dicke)
    linie.name = "Akzentlinie"
    return linie


def _punkt(slide, links: float, oben: float, groesse: float, farbe: str):
    from pptx.enum.shapes import MSO_SHAPE

    form = slide.shapes.add_shape(MSO_SHAPE.OVAL, _zoll(links), _zoll(oben), _zoll(groesse), _zoll(groesse))
    form.fill.solid()
    form.fill.fore_color.rgb = basis._rgb(farbe)
    form.line.fill.background()
    form.shadow.inherit = False
    form.name = "Deko-Punkt"
    return form


def _bogen(slide, links: float, oben: float, groesse: float, farbe: str, start: float = 270, ende: float = 20, dicke: float = 4.5, gestrichelt: bool = True):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    form = slide.shapes.add_shape(MSO_SHAPE.ARC, _zoll(links), _zoll(oben), _zoll(groesse), _zoll(groesse))
    form.adjustments[0] = start * 0.6
    form.adjustments[1] = ende * 0.6
    form.fill.background()
    form.line.color.rgb = basis._rgb(farbe)
    form.line.width = Pt(dicke)
    if gestrichelt:
        form.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    form.shadow.inherit = False
    form.name = "Deko-Bogen"
    return form


def _ring(slide, links: float, oben: float, groesse: float, farbe: str, dicke: float = 7):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    form = slide.shapes.add_shape(MSO_SHAPE.OVAL, _zoll(links), _zoll(oben), _zoll(groesse), _zoll(groesse))
    form.fill.background()
    form.line.color.rgb = basis._rgb(farbe)
    form.line.width = Pt(dicke)
    form.shadow.inherit = False
    form.name = "Deko-Ring"
    return form


def _foto(slide, bild: str, links: float, oben: float, breite: float, hoehe: float, rund: bool = False):
    from pptx.enum.shapes import MSO_SHAPE

    form = basis.einpassen(slide, bild, links, oben, breite, hoehe)
    if form is not None and rund:
        try:
            form.auto_shape_type = MSO_SHAPE.OVAL
        except Exception:
            pass
    return form


def _titel(slide, text: str, links: float, oben: float, breite: float, hoehe: float, farben: dict, gross: int = 40, klein: int = 28, farbe: str = "", ausrichtung: str = "left", anker: str = "bottom"):
    if not text:
        return 0.0
    groesse = gross
    while groesse > 20 and (not _wort_passt(text, breite, groesse) or (groesse > klein and not _passt([text], breite, groesse, hoehe))):
        groesse -= 2
    basis._text(slide, text, links, oben, breite, hoehe, groesse, farbe or farben["dark"], font=TITEL, align=ausrichtung, anchor=anker)
    return groesse


def _aufzaehlung(absatz, zeichen: str, einzug: float, farbe: str) -> None:
    from pptx.oxml.ns import qn

    eigenschaften = absatz._p.get_or_add_pPr()
    eigenschaften.set("marL", str(int(einzug * 914400)))
    eigenschaften.set("indent", str(int(-0.3 * 914400)))
    for alt in eigenschaften.findall(qn("a:buNone")) + eigenschaften.findall(qn("a:buChar")) + eigenschaften.findall(qn("a:buFont")) + eigenschaften.findall(qn("a:buClr")):
        eigenschaften.remove(alt)
    farbe_knoten = eigenschaften.makeelement(qn("a:buClr"), {})
    rgb = farbe_knoten.makeelement(qn("a:srgbClr"), {"val": farbe})
    farbe_knoten.append(rgb)
    eigenschaften.append(farbe_knoten)
    eigenschaften.append(eigenschaften.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    eigenschaften.append(eigenschaften.makeelement(qn("a:buChar"), {"char": zeichen}))


def _inhalt(slide, data: dict, links: float, oben: float, breite: float, hoehe: float, farben: dict, farbe: str = "") -> None:
    from pptx.util import Pt

    text = str(data.get("text", "") or "").strip()
    punkte = [str(p) for p in data.get("bullets") or [] if str(p).strip()]
    if not text and not punkte:
        return
    zeilen = ([text] if text else []) + punkte
    groesse = TEXT_GROSS
    while groesse > TEXT_KLEIN and not _passt(zeilen, breite - (0.6 if punkte else 0.3), groesse, hoehe):
        groesse -= 1
    box = slide.shapes.add_textbox(_zoll(links), _zoll(oben), _zoll(breite), _zoll(hoehe))
    rahmen = box.text_frame
    rahmen.word_wrap = True
    rahmen.margin_left = rahmen.margin_right = rahmen.margin_top = 0
    erster = True
    if text:
        absatz = rahmen.paragraphs[0]
        lauf = absatz.add_run()
        lauf.text = text
        lauf.font.size = Pt(groesse)
        lauf.font.name = TEXT
        lauf.font.color.rgb = basis._rgb(farbe or farben["dark"])
        absatz.space_after = Pt(10)
        absatz.line_spacing = 1.1
        erster = False
    for punkt in punkte:
        absatz = rahmen.paragraphs[0] if erster else rahmen.add_paragraph()
        erster = False
        unter = punkt.startswith(("  ", "– ", "- "))
        inhalt = punkt.lstrip(" –-").strip()
        absatz.space_after = Pt(4 if unter else 9)
        absatz.line_spacing = 1.08
        _aufzaehlung(absatz, "◦" if unter else "•", 0.65 if unter else 0.3, farbe or farben["dark"])
        kopf, trenner, rest = inhalt.partition(":")
        if trenner and 0 < len(kopf) <= 30:
            fett = absatz.add_run()
            fett.text = kopf + ":"
            fett.font.bold = True
            fett.font.size = Pt(groesse - (2 if unter else 0))
            fett.font.name = TEXT
            fett.font.color.rgb = basis._rgb(farbe or farben["dark"])
            inhalt = rest
        lauf = absatz.add_run()
        lauf.text = inhalt
        lauf.font.size = Pt(groesse - (2 if unter else 0))
        lauf.font.name = TEXT
        lauf.font.color.rgb = basis._rgb(farbe or farben["dark"])


def _weiss(slide, farben: dict) -> None:
    basis._fill(slide, farben["light"])


def titel_folie(prs, slide, data: dict, farben: dict) -> None:
    bild = str(data.get("image", "") or "")
    if bild and basis.einpassen(slide, bild, 0, 0, basis.SLIDE_W, basis.SLIDE_H) is not None:
        basis.schleier(slide, "000000", 46)
        schrift = farben["light"]
    else:
        basis._fill(slide, farben["tief"])
        _bogen(slide, basis.SLIDE_W - 2.7, -3.5, 6.0, farben["secondary"], 85, 195, 5)
        _punkt(slide, basis.SLIDE_W - 2.05, 1.38, 0.42, farben["accent"])
        _bogen(slide, -2.0, basis.SLIDE_H - 2.0, 4.0, farben["primary"], 265, 365, 4)
        schrift = farben["light"]
    titel = str(data.get("title", ""))
    _titel(slide, titel, 1.6, 1.2, basis.SLIDE_W - 3.2, 3.0, farben, 60, 40, schrift, "center")
    unter = str(data.get("subtitle", data.get("text", "")) or "")
    if unter:
        basis._text(slide, unter, 1.6, 4.55, basis.SLIDE_W - 3.2, 1.0, 20, schrift, font=TEXT, align="center")
    fuss = str(data.get("footer", "") or "")
    if fuss:
        basis._text(slide, fuss, 1.6, 6.6, basis.SLIDE_W - 3.2, 0.4, 13, schrift, font=TEXT, align="center")


def _menge(data: dict) -> int:
    punkte = [str(p) for p in data.get("bullets") or [] if str(p).strip()]
    return len(str(data.get("text", "") or "")) + sum(len(p) for p in punkte) + 25 * max(0, len(punkte) - 4)


def _variante(data: dict) -> int:
    return int(data.get("_nummer", 0)) % 3


def inhalt_folie(prs, slide, data: dict, farben: dict) -> None:
    _weiss(slide, farben)
    bild = str(data.get("image", "") or "")
    titel = str(data.get("title", "") or "")
    variante = _variante(data)
    if bild and variante == 1:
        if _foto(slide, bild, 7.2, 1.45, 6.9, 6.9, rund=True) is not None:
            _bogen(slide, 6.75, 1.0, 7.8, farben["primary"], 248, 330, 4.5)
            _punkt(slide, 7.05, 1.95, 0.5, farben["primary"])
            _titel(slide, titel, 0.9, 0.35, 5.9, 1.45, farben, 36, 26)
            _inhalt(slide, data, 0.9, 2.05, 5.9, 4.95, farben)
            return
    if bild and _menge(data) > 240:
        if _foto(slide, bild, 7.55, 0, basis.SLIDE_W - 7.55, basis.SLIDE_H) is not None:
            _titel(slide, titel, 0.66, 0.35, 6.4, 1.55, farben, 36, 26)
            _linie(slide, 0.66, 2.05, 3.7, farben["accent"])
            _inhalt(slide, data, 0.66, 2.4, 6.4, 4.7, farben)
            return
    if bild and _foto(slide, bild, 5.85, 0, basis.SLIDE_W - 5.85, basis.SLIDE_H) is not None:
        _titel(slide, titel, 0.66, 0.35, 4.85, 2.05, farben, 40, 28)
        _linie(slide, 0.66, 2.62, 3.7, farben["accent"])
        _inhalt(slide, data, 0.66, 3.0, 4.75, 3.95, farben)
        return
    _titel(slide, titel, 0.9, 0.45, basis.SLIDE_W - 1.8, 1.35, farben, 40, 28)
    _linie(slide, 0.9, 1.98, 3.7, farben["accent"])
    _ring(slide, basis.SLIDE_W - 3.1, basis.SLIDE_H - 3.0, 4.4, farben["soft"], 10)
    _inhalt(slide, data, 0.9, 2.35, basis.SLIDE_W - 3.4, 4.6, farben)


def bild_folie(prs, slide, data: dict, farben: dict) -> None:
    bild = str(data.get("image", "") or "")
    if not bild:
        inhalt_folie(prs, slide, data, farben)
        return
    _weiss(slide, farben)
    if _foto(slide, bild, 0, 0, 7.4, basis.SLIDE_H) is None:
        inhalt_folie(prs, slide, data, farben)
        return
    _titel(slide, str(data.get("title", "") or ""), 8.0, 0.5, 4.7, 1.9, farben, 36, 26)
    _linie(slide, 8.0, 2.62, 3.2, farben["accent"])
    _inhalt(slide, data, 8.0, 3.0, 4.75, 4.0, farben)


def dunkel_folie(prs, slide, data: dict, farben: dict) -> None:
    bild = str(data.get("image", "") or "")
    if bild and basis.einpassen(slide, bild, 0, 0, basis.SLIDE_W, basis.SLIDE_H) is not None:
        basis.schleier(slide, "000000", 66)
    else:
        basis._fill(slide, "0B0B0F")
        _bogen(slide, 7.8, 0.4, 5.2, farben["secondary"], 200, 40, 3)
    hell = farben["light"]
    titel = str(data.get("title", "") or "")
    if data.get("bullets") or data.get("items"):
        punkte = list(data.get("bullets") or [])
        for eintrag in data.get("items") or []:
            if isinstance(eintrag, dict):
                punkte.append(str(eintrag.get("titel") or eintrag.get("title") or "") + ":")
                punkte.extend("  " + str(b) for b in eintrag.get("bullets") or [])
                if eintrag.get("text"):
                    punkte.append("  " + str(eintrag["text"]))
        _titel(slide, titel, 0.9, 0.6, 4.9, 6.0, farben, 40, 28, hell, anker="middle")
        _inhalt(slide, {"bullets": punkte, "text": data.get("text", "")}, 6.6, 0.6, 5.9, 6.3, farben, hell)
        return
    zitat = str(data.get("text", "") or titel)
    _titel(slide, f"„{zitat}“" if data.get("layout") == "quote" else zitat, 1.4, 1.6, basis.SLIDE_W - 2.8, 3.6, farben, 40, 26, hell, "center", "middle")
    quelle = str(data.get("subtitle", data.get("source", "")) or "")
    if quelle:
        basis._text(slide, f"— {quelle}", 1.4, 5.5, basis.SLIDE_W - 2.8, 0.6, 18, farben["accent"], font=TEXT, align="center")


def text_folie(prs, slide, data: dict, farben: dict) -> None:
    roh = data.get("absaetze") or data.get("paragraphs") or []
    if roh and not data.get("text"):
        data = {**data, "text": "\n\n".join(str(t) for t in roh if str(t).strip())}
    inhalt_folie(prs, slide, data, farben)


def agenda_folie(prs, slide, data: dict, farben: dict) -> None:
    from app.services.pptx_inhalte import punkt_text

    _weiss(slide, farben)
    bild = str(data.get("image", "") or "")
    breite = basis.SLIDE_W - 1.8
    if bild and _foto(slide, bild, 8.4, 0, basis.SLIDE_W - 8.4, basis.SLIDE_H) is not None:
        breite = 6.9
    _titel(slide, str(data.get("title", "Agenda") or "Agenda"), 0.9, 0.45, breite, 1.35, farben, 40, 28)
    _linie(slide, 0.9, 1.98, 3.7, farben["accent"])
    eintraege = (data.get("items") or data.get("bullets") or [])[:7]
    if not eintraege:
        return
    platz = basis.SLIDE_H - 2.55 - 0.5
    hoehe = min(0.95, platz / len(eintraege))
    for stelle, eintrag in enumerate(eintraege):
        kopf, rest = punkt_text(eintrag)
        oben = 2.45 + stelle * hoehe
        basis._text(slide, f"{stelle + 1:02d}", 0.9, oben, 0.8, 0.5, 26, farben["accent"], font=TITEL)
        basis._text(slide, kopf, 1.75, oben + 0.02, breite - 0.9, 0.46, 22, farben["dark"], font=TEXT)
        if rest and hoehe >= 0.75:
            basis._text(slide, rest, 1.75, oben + 0.46, breite - 0.9, hoehe - 0.46, 15, farben["muted"], font=TEXT)


def schluss_folie(prs, slide, data: dict, farben: dict) -> None:
    _weiss(slide, farben)
    titel = str(data.get("title", "") or "Vielen Dank für Ihre Aufmerksamkeit")
    bild = str(data.get("image", "") or "")
    if bild and _foto(slide, bild, 7.0, 0.8, 5.9, 5.9, rund=True) is not None:
        _bogen(slide, 6.4, 0.3, 3.0, farben["primary"])
    else:
        _ring(slide, 7.25, 0.85, 5.6, "D9D9D9", 9)
        _punkt(slide, 8.85, 2.55, 0.62, "D9D9D9")
        _punkt(slide, 10.65, 2.55, 0.62, "D9D9D9")
        _bogen(slide, 8.45, 2.05, 3.55, "D9D9D9", 25, 155, 9, False)
    _ring(slide, 0.9, 3.05, 0.95, farben["primary"], 4)
    _punkt(slide, 1.17, 3.33, 0.13, farben["primary"])
    _punkt(slide, 1.5, 3.33, 0.13, farben["primary"])
    _bogen(slide, 1.12, 3.27, 0.52, farben["primary"], 25, 155, 3, False)
    _titel(slide, titel, 2.15, 1.3, 4.8, 4.4, farben, 48, 32, anker="middle")
    unter = str(data.get("subtitle", data.get("text", "")) or "")
    if unter:
        basis._text(slide, unter, 2.15, 5.75, 4.8, 1.0, 16, farben["muted"], font=TEXT)


LAYOUTS = {
    "title": titel_folie,
    "bullets": inhalt_folie,
    "image": bild_folie,
    "quote": dunkel_folie,
    "dunkel": dunkel_folie,
    "closing": schluss_folie,
    "text": text_folie,
    "prosa": text_folie,
    "agenda": agenda_folie,
}
