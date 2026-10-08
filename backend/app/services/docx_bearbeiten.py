from __future__ import annotations

import asyncio
import copy
from pathlib import Path

STILE = {
    "titel": "Title",
    "h1": "Heading 1",
    "h2": "Heading 2",
    "h3": "Heading 3",
    "h4": "Heading 4",
    "liste": "List Bullet",
    "nummer": "List Number",
    "zitat": "Quote",
    "absatz": "Normal",
}
BILDFELDER = ("image", "bild", "image_url", "bild_url", "bild_suche", "image_search", "bild_prompt", "image_prompt")


def _stil(dokument, name: str):
    wunsch = STILE.get(str(name or "absatz").strip().lower(), str(name or "Normal"))
    try:
        return dokument.styles[wunsch]
    except KeyError:
        return None


def _absatz_text(absatz, text: str) -> None:
    laeufe = absatz.runs
    if not laeufe:
        absatz.add_run(text)
        return
    laeufe[0].text = text
    for rest in laeufe[1:]:
        rest._r.getparent().remove(rest._r)


def _neuer_absatz(dokument, anker, text: str, stil: str, davor: bool = False):
    absatz = dokument.add_paragraph(text)
    gefunden = _stil(dokument, stil)
    if gefunden is not None:
        absatz.style = gefunden
    if anker is not None:
        (anker.addprevious if davor else anker.addnext)(absatz._p)
    return absatz


async def bearbeiten(pfad: str, aktionen: list[dict]) -> dict:
    from app.services.bildquelle import besorgen

    quelle = Path(str(pfad or "")).expanduser()
    if not quelle.is_file() or quelle.suffix.lower() != ".docx":
        return {"error": f"Word-Datei nicht gefunden: {quelle}"}
    if not isinstance(aktionen, list) or not aktionen:
        return {"error": "aktionen fehlt. Gib eine Liste von Bearbeitungsschritten an."}
    for aktion in aktionen:
        if isinstance(aktion, dict) and str(aktion.get("aktion", "")).lower() == "bild" and any(str(aktion.get(f) or "").strip() for f in BILDFELDER):
            aktion["_bild"] = await besorgen(aktion, True)
    return await asyncio.to_thread(_anwenden, quelle, aktionen)


def _anwenden(quelle: Path, aktionen: list[dict]) -> dict:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph

    from app.services.dateiansicht_service import docx_elemente

    dokument = Document(str(quelle))
    elemente = docx_elemente(dokument)
    erledigt: list[str] = []
    fehler: list[str] = []

    def element(nummer) -> object:
        index = int(nummer)
        if index < 0 or index >= len(elemente):
            raise ValueError(f"Absatz {index} gibt es nicht (0 bis {len(elemente) - 1}).")
        return elemente[index]

    def absatz(nummer) -> Paragraph:
        kind = element(nummer)
        if not kind.tag.endswith("}p"):
            raise ValueError(f"Block {nummer} ist eine Tabelle, kein Absatz.")
        return Paragraph(kind, dokument._body)

    def anker(aktion: dict):
        nach = aktion.get("nach")
        if nach is None:
            return (elemente[-1] if elemente else None), False
        if int(nach) < 0:
            return (elemente[0] if elemente else None), True
        return element(nach), False

    for schritt, aktion in enumerate(aktionen, start=1):
        if not isinstance(aktion, dict):
            fehler.append(f"Schritt {schritt}: kein Objekt")
            continue
        art = str(aktion.get("aktion", "")).strip().lower()
        try:
            if art == "text":
                ziel = absatz(aktion.get("nr", aktion.get("block")))
                _absatz_text(ziel, str(aktion.get("text", "")))
                if aktion.get("stil"):
                    gefunden = _stil(dokument, str(aktion["stil"]))
                    if gefunden is not None:
                        ziel.style = gefunden
                erledigt.append(f"Absatz {aktion.get('nr')} geändert")
            elif art == "ersetzen":
                suchen, ersetzen = str(aktion.get("suchen", "")), str(aktion.get("ersetzen", ""))
                if not suchen:
                    raise ValueError("suchen fehlt.")
                anzahl = 0
                alle = list(dokument.paragraphs) + [p for t in dokument.tables for r in t.rows for c in r.cells for p in c.paragraphs]
                for eintrag in alle:
                    if suchen in eintrag.text:
                        anzahl += eintrag.text.count(suchen)
                        treffer = [l for l in eintrag.runs if suchen in l.text]
                        if treffer:
                            for lauf in treffer:
                                lauf.text = lauf.text.replace(suchen, ersetzen)
                        else:
                            _absatz_text(eintrag, eintrag.text.replace(suchen, ersetzen))
                erledigt.append(f"„{suchen}“ {anzahl}× ersetzt")
            elif art == "einfuegen":
                ort, davor = anker(aktion)
                zeilen = [z for z in str(aktion.get("text", "")).split("\n")]
                stil = str(aktion.get("stil", "absatz"))
                for zeile in (reversed(zeilen) if not davor else zeilen):
                    _neuer_absatz(dokument, ort, zeile, stil, davor)
                erledigt.append(f"{len(zeilen)} Absatz/Absätze eingefügt")
            elif art == "loeschen":
                kind = element(aktion.get("nr", aktion.get("block")))
                kind.getparent().remove(kind)
                erledigt.append(f"Block {aktion.get('nr')} gelöscht")
            elif art == "format":
                ziel = absatz(aktion.get("nr", aktion.get("block")))
                if aktion.get("stil"):
                    gefunden = _stil(dokument, str(aktion["stil"]))
                    if gefunden is not None:
                        ziel.style = gefunden
                if aktion.get("ausrichtung"):
                    ziel.alignment = {"mitte": WD_ALIGN_PARAGRAPH.CENTER, "center": WD_ALIGN_PARAGRAPH.CENTER, "rechts": WD_ALIGN_PARAGRAPH.RIGHT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "block": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(str(aktion["ausrichtung"]).lower(), WD_ALIGN_PARAGRAPH.LEFT)
                for lauf in ziel.runs:
                    if "fett" in aktion:
                        lauf.bold = bool(aktion["fett"])
                    if "kursiv" in aktion:
                        lauf.italic = bool(aktion["kursiv"])
                    if aktion.get("groesse"):
                        lauf.font.size = Pt(float(aktion["groesse"]))
                erledigt.append(f"Absatz {aktion.get('nr')} formatiert")
            elif art == "bild":
                bild = str(aktion.get("_bild") or "")
                if not bild or not Path(bild).is_file():
                    raise ValueError("Kein Bild gefunden. Gib image (Pfad oder URL), bild_suche oder bild_prompt an.")
                ort, davor = anker(aktion)
                neu = _neuer_absatz(dokument, ort, "", "absatz", davor)
                neu.add_run().add_picture(bild, width=Cm(float(aktion.get("breite_cm") or 15)))
                neu.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if aktion.get("unterschrift"):
                    unter = _neuer_absatz(dokument, neu._p, str(aktion["unterschrift"]), "absatz")
                    unter.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for lauf in unter.runs:
                        lauf.italic = True
                erledigt.append("Bild eingefügt")
            elif art == "tabelle":
                zeilen = [[str(z) for z in reihe] for reihe in aktion.get("zeilen") or [] if isinstance(reihe, list)]
                if not zeilen:
                    raise ValueError("zeilen fehlt.")
                spalten = max(len(z) for z in zeilen)
                tabelle = dokument.add_table(rows=len(zeilen), cols=spalten)
                try:
                    tabelle.style = dokument.styles["Table Grid"]
                except KeyError:
                    pass
                for r, reihe in enumerate(zeilen):
                    for s, wert in enumerate(reihe):
                        tabelle.cell(r, s).text = wert
                for zelle in tabelle.rows[0].cells:
                    for lauf in (zelle.paragraphs[0].runs if zelle.paragraphs else []):
                        lauf.bold = True
                ort, davor = anker(aktion)
                if ort is not None:
                    (ort.addprevious if davor else ort.addnext)(tabelle._tbl)
                erledigt.append(f"Tabelle mit {len(zeilen)} Zeilen eingefügt")
            elif art == "zelle":
                kind = element(aktion.get("nr", aktion.get("block")))
                from docx.table import Table

                tabelle = Table(kind, dokument._body)
                tabelle.cell(int(aktion.get("zeile", 0)), int(aktion.get("spalte", 0))).text = str(aktion.get("text", ""))
                erledigt.append("Tabellenzelle geändert")
            elif art == "duplizieren":
                kind = element(aktion.get("nr", aktion.get("block")))
                kind.addnext(copy.deepcopy(kind))
                erledigt.append(f"Block {aktion.get('nr')} dupliziert")
            else:
                raise ValueError(f"Unbekannte aktion '{art}'.")
        except Exception as exc:
            fehler.append(f"Schritt {schritt} ({art or '?'}): {exc}")
    if not erledigt:
        return {"error": "Nichts geändert. " + " | ".join(fehler)}
    ziel = quelle
    try:
        dokument.save(str(quelle))
    except PermissionError:
        ziel = quelle.with_name(f"{quelle.stem} (bearbeitet).docx")
        dokument.save(str(ziel))
    ergebnis: dict = {"ok": True, "path": str(ziel), "erledigt": erledigt}
    if ziel != quelle:
        ergebnis["hinweis"] = "Die Originaldatei ist gerade in Word geöffnet und gesperrt - die Änderungen liegen in einer Kopie."
    if fehler:
        ergebnis["probleme"] = fehler
    return ergebnis


def lesen(pfad: str) -> dict:
    from app.services.dateiansicht_service import docx_bloecke
    from app.services.kommentar_service import beschreibung, get_kommentar_service

    quelle = Path(str(pfad or "")).expanduser()
    if not quelle.is_file():
        return {"error": f"Datei nicht gefunden: {quelle}"}
    bloecke = []
    for block in docx_bloecke(quelle):
        eintrag = {"nr": block["nr"], "art": block["art"]}
        if block.get("text"):
            eintrag["text"] = block["text"][:1500]
        if block.get("zeilen"):
            eintrag["zeilen"] = block["zeilen"][:30]
        if block.get("bilder"):
            eintrag["bilder"] = len(block["bilder"])
        if len(eintrag) > 2:
            bloecke.append(eintrag)
    ergebnis = {"path": str(quelle), "bloecke": bloecke}
    kommentare = [
        {"id": k["id"], "wo": beschreibung(k.get("stelle") or {}), "stelle": k.get("stelle"), "text": k["text"]}
        for k in get_kommentar_service().liste(quelle, offen=True)
    ]
    if kommentare:
        ergebnis["offene_kommentare"] = kommentare
        ergebnis["hinweis"] = "Der Nutzer hat Kommentare hinterlassen. Setze sie mit edit_docx um und markiere sie danach mit datei_kommentare (aktion erledigt) als erledigt."
    return ergebnis
