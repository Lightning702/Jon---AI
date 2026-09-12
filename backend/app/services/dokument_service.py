from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from app.core.fehler import leise
from app.services.dateiindex_service import get_dateiindex_service, karte
from app.services.dateiraum_service import get_dateiraum_service, saeubern

TEXTARTEN = {
    "txt",
    "md",
    "markdown",
    "json",
    "yaml",
    "yml",
    "xml",
    "html",
    "css",
    "js",
    "ts",
    "tsx",
    "jsx",
    "py",
    "sql",
    "c",
    "cpp",
    "h",
    "hpp",
    "java",
    "rs",
    "go",
    "dart",
    "sh",
    "ps1",
    "bat",
    "ini",
    "toml",
    "env",
    "log",
    "rtf",
}

BEKANNT = {"pdf", "docx", "odt", "xlsx", "ods", "csv"} | TEXTARTEN

UEBERSCHRIFT = re.compile(r"^(#{1,4})\s+(.*)$")
PUNKT = re.compile(r"^\s*[-*•]\s+(.*)$")
NUMMER = re.compile(r"^\s*(\d{1,2})[.)]\s+(.*)$")

FEHLT = {
    "pdf": "reportlab",
    "docx": "python-docx",
    "odt": "odfpy",
    "ods": "odfpy",
    "xlsx": "openpyxl",
}


def _bloecke(text: str, titel: str = "") -> list[tuple[str, str]]:
    bloecke: list[tuple[str, str]] = []
    kopf = _vergleichbar(titel)
    for zeile in str(text or "").replace("\r\n", "\n").split("\n"):
        roh = zeile.rstrip()
        if not roh.strip():
            bloecke.append(("leer", ""))
            continue
        treffer = UEBERSCHRIFT.match(roh)
        if treffer:
            text_h = treffer.group(2).strip()
            stufe = len(treffer.group(1))
            if stufe == 1 and kopf and _vergleichbar(text_h) == kopf:
                continue
            bloecke.append((f"h{stufe}", text_h))
            continue
        treffer = PUNKT.match(roh)
        if treffer:
            bloecke.append(("punkt", treffer.group(1).strip()))
            continue
        treffer = NUMMER.match(roh)
        if treffer:
            bloecke.append(("nummer", treffer.group(2).strip()))
            continue
        bloecke.append(("absatz", roh.strip()))
    return bloecke


def _vergleichbar(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def _ohne_auszeichnung(text: str) -> str:
    ohne = re.sub(r"\*\*(.+?)\*\*", r"\1", str(text or ""))
    ohne = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", ohne)
    ohne = re.sub(r"`(.+?)`", r"\1", ohne)
    return ohne


def _fett_xml(text: str) -> str:
    from xml.sax.saxutils import escape

    sicher = escape(_ohne_auszeichnung(text))
    return sicher


def _zeilen_aus(inhalt) -> list[list[str]]:
    if isinstance(inhalt, list):
        zeilen = []
        for eintrag in inhalt:
            if isinstance(eintrag, (list, tuple)):
                zeilen.append([str(w) for w in eintrag])
            elif isinstance(eintrag, dict):
                if not zeilen:
                    zeilen.append([str(k) for k in eintrag.keys()])
                zeilen.append([str(w) for w in eintrag.values()])
            else:
                zeilen.append([str(eintrag)])
        return zeilen
    text = str(inhalt or "").strip()
    if not text:
        return []
    try:
        daten = json.loads(text)
        if isinstance(daten, list):
            return _zeilen_aus(daten)
    except Exception as fehler:
        leise(fehler, "services/dokument_service")
    trenner = ";" if text.count(";") > text.count(",") else ","
    return [
        [feld.strip() for feld in zeile.split(trenner)]
        for zeile in text.split("\n")
        if zeile.strip()
    ]


class DokumentService:
    def _ziel(self, art: str, titel: str, wunsch: str, dateiname: str) -> dict:
        endung = art.lower().lstrip(".")
        if endung == "markdown":
            endung = "md"
        name = saeubern(dateiname or f"{titel or 'Dokument'}.{endung}")
        if not Path(name).suffix:
            name = f"{name}.{endung}"
        return get_dateiraum_service().zielpfad(name, wunsch)

    def erstellen(
        self,
        art: str,
        titel: str = "",
        inhalt: str = "",
        wunsch: str = "",
        dateiname: str = "",
        projekt: str = "",
        auftrag: str = "",
        quelle: str = "app",
    ) -> dict:
        endung = str(art or "").lower().lstrip(".").strip()
        if endung == "markdown":
            endung = "md"
        if endung == "text":
            endung = "txt"
        if not endung:
            endung = "txt"
        if endung not in BEKANNT:
            return {
                "error": f"Dateityp .{endung} kann ich noch nicht erzeugen.",
                "moeglich": sorted(BEKANNT),
            }
        ziel = self._ziel(endung, titel, wunsch, dateiname)
        if ziel.get("error"):
            return ziel
        pfad = Path(ziel["pfad"])
        bauer = {
            "pdf": self._pdf,
            "docx": self._docx,
            "odt": self._odt,
            "xlsx": self._xlsx,
            "ods": self._ods,
            "csv": self._csv,
        }.get(endung, self._text)
        try:
            bauer(pfad, titel, inhalt)
        except ModuleNotFoundError as exc:
            paket = FEHLT.get(endung, str(exc).split("'")[-2] if "'" in str(exc) else "")
            return {
                "error": (
                    f"Fuer .{endung} fehlt die Bibliothek {paket}. "
                    f"Installiere sie mit: pip install {paket}"
                )
            }
        except Exception as exc:
            return {"error": f"Die Datei liess sich nicht erzeugen: {exc}"}
        if not pfad.exists():
            return {"error": "Die Datei wurde nicht geschrieben."}
        beschreibung = _ohne_auszeichnung(str(inhalt or ""))[:400]
        get_dateiindex_service().merken(
            pfad, titel or pfad.stem, beschreibung, projekt, auftrag, quelle
        )
        return {"ok": True, "datei": karte(pfad, projekt, titel or pfad.stem)}

    @staticmethod
    def _text(pfad: Path, titel: str, inhalt: str) -> None:
        text = str(inhalt or "")
        if titel and pfad.suffix.lower() in (".txt", ".md"):
            kopf = f"# {titel}\n\n" if pfad.suffix.lower() == ".md" else f"{titel}\n\n"
            if not text.lstrip().startswith(("#", titel)):
                text = kopf + text
        pfad.write_text(text, encoding="utf-8")

    @staticmethod
    def _csv(pfad: Path, titel: str, inhalt: str) -> None:
        zeilen = _zeilen_aus(inhalt)
        with pfad.open("w", encoding="utf-8-sig", newline="") as datei:
            schreiber = csv.writer(datei, delimiter=";")
            for zeile in zeilen:
                schreiber.writerow(zeile)

    @staticmethod
    def _pdf(pfad: Path, titel: str, inhalt: str) -> None:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )

        vorlagen = getSampleStyleSheet()
        stil_titel = ParagraphStyle(
            "JonTitel",
            parent=vorlagen["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=27,
            spaceAfter=14,
            alignment=TA_LEFT,
        )
        stil_h = {
            1: ParagraphStyle(
                "JonH1",
                parent=vorlagen["Heading1"],
                fontName="Helvetica-Bold",
                fontSize=16,
                leading=20,
                spaceBefore=14,
                spaceAfter=6,
            ),
            2: ParagraphStyle(
                "JonH2",
                parent=vorlagen["Heading2"],
                fontName="Helvetica-Bold",
                fontSize=13,
                leading=17,
                spaceBefore=11,
                spaceAfter=5,
            ),
        }
        stil_h[3] = stil_h[2]
        stil_h[4] = stil_h[2]
        stil_text = ParagraphStyle(
            "JonText",
            parent=vorlagen["BodyText"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=15.5,
            spaceAfter=7,
        )
        dokument = SimpleDocTemplate(
            str(pfad),
            pagesize=A4,
            leftMargin=22 * mm,
            rightMargin=22 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm,
            title=titel or pfad.stem,
            author="Jon",
        )
        fluss = []
        if titel:
            fluss.append(Paragraph(_fett_xml(titel), stil_titel))
        punkte: list[str] = []
        punkt_art = ""

        def punkte_leeren() -> None:
            nonlocal punkt_art
            if not punkte:
                return
            nummeriert = punkt_art == "nummer"
            fluss.append(
                ListFlowable(
                    [
                        ListItem(Paragraph(_fett_xml(p), stil_text), leftIndent=14)
                        for p in punkte
                    ],
                    bulletType="1" if nummeriert else "bullet",
                    bulletFormat="%s." if nummeriert else None,
                    start=1 if nummeriert else "circle",
                    bulletFontSize=8 if not nummeriert else 10.5,
                    leftIndent=16,
                )
            )
            fluss.append(Spacer(1, 5))
            punkte.clear()
            punkt_art = ""

        for art, text in _bloecke(inhalt, titel):
            if art in ("punkt", "nummer"):
                if punkt_art and punkt_art != art:
                    punkte_leeren()
                punkt_art = art
                punkte.append(text)
                continue
            punkte_leeren()
            if art == "leer":
                continue
            if art.startswith("h"):
                stufe = int(art[1])
                if stufe == 1 and fluss:
                    fluss.append(Spacer(1, 6))
                fluss.append(Paragraph(_fett_xml(text), stil_h.get(stufe, stil_h[2])))
                continue
            if text == "---":
                fluss.append(PageBreak())
                continue
            fluss.append(Paragraph(_fett_xml(text), stil_text))
        punkte_leeren()
        if not fluss:
            fluss.append(Paragraph("(leer)", stil_text))
        dokument.build(fluss)

    @staticmethod
    def _docx(pfad: Path, titel: str, inhalt: str) -> None:
        from docx import Document
        from docx.shared import Pt

        dokument = Document()
        stil = dokument.styles["Normal"]
        stil.font.name = "Calibri"
        stil.font.size = Pt(11)
        if titel:
            dokument.add_heading(_ohne_auszeichnung(titel), level=0)
        for art, text in _bloecke(inhalt, titel):
            sauber = _ohne_auszeichnung(text)
            if art == "leer":
                continue
            if art.startswith("h"):
                dokument.add_heading(sauber, level=min(4, int(art[1])))
            elif art == "punkt":
                dokument.add_paragraph(sauber, style="List Bullet")
            elif art == "nummer":
                dokument.add_paragraph(sauber, style="List Number")
            else:
                dokument.add_paragraph(sauber)
        dokument.save(str(pfad))

    @staticmethod
    def _odt(pfad: Path, titel: str, inhalt: str) -> None:
        from odf.opendocument import OpenDocumentText
        from odf.style import ParagraphProperties, Style, TextProperties
        from odf.text import H, List, ListItem, P

        dokument = OpenDocumentText()
        stil = Style(name="JonText", family="paragraph")
        stil.addElement(ParagraphProperties(marginbottom="0.15cm"))
        stil.addElement(TextProperties(fontsize="11pt"))
        dokument.automaticstyles.addElement(stil)
        if titel:
            dokument.text.addElement(H(outlinelevel=1, text=_ohne_auszeichnung(titel)))
        liste = None
        for art, text in _bloecke(inhalt, titel):
            sauber = _ohne_auszeichnung(text)
            if art in ("punkt", "nummer"):
                if liste is None:
                    liste = List()
                    dokument.text.addElement(liste)
                eintrag = ListItem()
                eintrag.addElement(P(stylename=stil, text=sauber))
                liste.addElement(eintrag)
                continue
            liste = None
            if art == "leer":
                continue
            if art.startswith("h"):
                dokument.text.addElement(
                    H(outlinelevel=min(4, int(art[1]) + 1), text=sauber)
                )
                continue
            dokument.text.addElement(P(stylename=stil, text=sauber))
        dokument.save(str(pfad))

    @staticmethod
    def _xlsx(pfad: Path, titel: str, inhalt: str) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        zeilen = _zeilen_aus(inhalt)
        mappe = Workbook()
        blatt = mappe.active
        blatt.title = saeubern(titel or "Tabelle", "Tabelle")[:31] or "Tabelle"
        for zeile in zeilen:
            blatt.append(zeile)
        if zeilen:
            kopf = PatternFill("solid", start_color="FFD4AF37", end_color="FFD4AF37")
            for zelle in blatt[1]:
                zelle.font = Font(bold=True)
                zelle.fill = kopf
                zelle.alignment = Alignment(vertical="center")
            blatt.freeze_panes = "A2"
            for spalte in range(1, (max(len(z) for z in zeilen) or 1) + 1):
                laenge = max(
                    (len(str(z[spalte - 1])) for z in zeilen if len(z) >= spalte),
                    default=8,
                )
                blatt.column_dimensions[get_column_letter(spalte)].width = min(
                    60, max(10, laenge + 2)
                )
        mappe.save(str(pfad))

    @staticmethod
    def _ods(pfad: Path, titel: str, inhalt: str) -> None:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        zeilen = _zeilen_aus(inhalt)
        dokument = OpenDocumentSpreadsheet()
        tabelle = Table(name=saeubern(titel or "Tabelle", "Tabelle")[:31] or "Tabelle")
        for zeile in zeilen:
            reihe = TableRow()
            for wert in zeile:
                zelle = TableCell(valuetype="string")
                zelle.addElement(P(text=str(wert)))
                reihe.addElement(zelle)
            tabelle.addElement(reihe)
        dokument.spreadsheet.addElement(tabelle)
        dokument.save(str(pfad))


_service: DokumentService | None = None


def get_dokument_service() -> DokumentService:
    global _service
    if _service is None:
        _service = DokumentService()
    return _service
