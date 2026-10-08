from __future__ import annotations

import asyncio
import base64
import io
from pathlib import Path

from app.core.config import DATA_DIR, get_settings
from app.core.fehler import leise
from app.providers.openai_compatible import OpenAICompatibleProvider
from app.providers.registry import get_registry
from app.services.screen_service import VISION_DEFAULTS
from app.services.settings_service import get_settings_service

MAX_BYTES = 300_000_000
MAX_TEXT = 20_000
ANHANG_ORDNER = DATA_DIR / "anhaenge"
BILD_KANTE = 1568

OFFICE = {"pptx", "docx", "xlsx", "xlsm", "odt", "ods", "odp", "rtf"}
BILDER = {"png", "jpg", "jpeg", "webp", "gif", "bmp", "heic", "tif", "tiff"}

IMAGE_PROMPT = (
    "Der Nutzer hat dieses Bild in den Chat gezogen. Beschreibe praezise und "
    "detailliert auf Deutsch, was darauf zu sehen ist: Inhalt, Text (woertlich "
    "zitieren, falls lesbar), Zahlen, Diagramme, Fehlermeldungen, Layout. "
    "Keine Einleitung, direkt die Beschreibung."
)


def _sauber(name: str) -> str:
    return "".join(
        zeichen if zeichen.isalnum() or zeichen in "._- " else "_"
        for zeichen in Path(str(name or "")).name.strip()
    ).strip(" .")[:120] or "anhang"


def upload_ordner() -> Path:
    try:
        from app.services.dateiraum_service import get_dateiraum_service

        ordner = get_dateiraum_service().sicherstellen() / "Uploads"
        ordner.mkdir(parents=True, exist_ok=True)
        return ordner
    except Exception as fehler:
        leise(fehler, "services/attachment_service")
        ANHANG_ORDNER.mkdir(parents=True, exist_ok=True)
        return ANHANG_ORDNER


def _ziel(name: str) -> Path:
    ordner = upload_ordner()
    sauber = _sauber(name)
    ziel = ordner / sauber
    if not ziel.exists():
        return ziel
    stamm, endung = Path(sauber).stem, Path(sauber).suffix
    zaehler = 2
    while True:
        ziel = ordner / f"{stamm} ({zaehler}){endung}"
        if not ziel.exists():
            return ziel
        zaehler += 1


def _endung(name: str) -> str:
    return Path(str(name or "")).suffix.lower().lstrip(".")


def _art(name: str, mime: str) -> str:
    endung = _endung(name)
    if mime.startswith("image/") or endung in BILDER:
        return "image"
    if mime == "application/pdf" or endung == "pdf":
        return "pdf"
    if endung in ("pptx", "odp"):
        return "praesentation"
    if endung in ("docx", "odt", "rtf"):
        return "dokument"
    if endung in ("xlsx", "xlsm", "ods"):
        return "tabelle"
    return "text"


class AttachmentService:
    async def extract(
        self,
        name: str,
        mime: str,
        data_b64: str,
        provider_name: str | None = None,
    ) -> dict:
        try:
            raw = base64.b64decode(data_b64)
        except Exception:
            return {"error": "Ungueltige Daten"}
        if len(raw) > MAX_BYTES:
            return {"error": "Datei zu gross (max 300 MB)"}
        ziel = _ziel(name)
        try:
            await asyncio.to_thread(ziel.write_bytes, raw)
        except OSError as fehler:
            return {"error": f"Datei konnte nicht gespeichert werden: {fehler}"}
        return await self.auswerten(ziel, name, mime, provider_name, True)

    async def hochladen(self, name: str, mime: str, stream, provider_name: str | None = None) -> dict:
        ziel = _ziel(name)
        groesse = 0
        try:
            with ziel.open("xb") as datei:
                async for stueck in stream:
                    groesse += len(stueck)
                    if groesse > MAX_BYTES:
                        raise ValueError("Datei zu gross (max 300 MB)")
                    datei.write(stueck)
            if not groesse:
                raise ValueError("Die Datei ist leer.")
        except BaseException:
            ziel.unlink(missing_ok=True)
            raise
        return await self.auswerten(ziel, name, mime, provider_name)

    async def auswerten(self, pfad: Path, name: str, mime: str, provider_name: str | None = None, beschreiben: bool = False) -> dict:
        art = _art(name, mime)
        if art == "image":
            ergebnis = await self._describe_image(name, mime, pfad, provider_name) if beschreiben else _bildhinweis(name, pfad)
        else:
            ergebnis = await asyncio.to_thread(self._lesen, art, name, pfad)
        ergebnis["pfad"] = str(pfad)
        ergebnis["groesse"] = pfad.stat().st_size if pfad.exists() else 0
        if "error" in ergebnis:
            ergebnis = {"kind": art, "name": name, "pfad": str(pfad), "content": f"Die Datei wurde gespeichert, ihr Inhalt ließ sich aber nicht auslesen ({ergebnis['error']}).", "groesse": ergebnis["groesse"]}
        try:
            from app.services.dateiindex_service import get_dateiindex_service

            ergebnis["datei"] = get_dateiindex_service().karte_und_merken(str(pfad), Path(name).stem, "Vom Nutzer hochgeladen", quelle="upload")
        except Exception as fehler:
            leise(fehler, "services/attachment_service")
        return ergebnis

    def _lesen(self, art: str, name: str, pfad: Path) -> dict:
        if art == "pdf":
            return self._extract_pdf(name, pfad)
        if art == "praesentation":
            return self._extract_praesentation(name, pfad)
        if art in ("dokument", "tabelle"):
            return self._extract_office(name, pfad, art)
        return self._extract_text(name, pfad)

    def _extract_pdf(self, name: str, pfad: Path) -> dict:
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(pfad))
            total = len(reader.pages)
            teile, laenge = [], 0
            for seite in reader.pages[:60]:
                text = seite.extract_text() or ""
                teile.append(text)
                laenge += len(text)
                if laenge > MAX_TEXT:
                    break
        except Exception as exc:
            return {"error": f"PDF konnte nicht gelesen werden: {exc}"}
        return {"kind": "pdf", "name": name, "pages": total, "content": "\n\n".join(teile)[:MAX_TEXT]}

    def _extract_praesentation(self, name: str, pfad: Path) -> dict:
        if _endung(name) == "odp":
            return self._extract_odf(name, pfad, "praesentation")
        from pptx import Presentation

        try:
            praesentation = Presentation(str(pfad))
        except Exception as exc:
            return {"error": f"PowerPoint konnte nicht gelesen werden: {exc}"}
        zeilen = []
        for nummer, folie in enumerate(praesentation.slides, start=1):
            teile = []
            for index, form in enumerate(folie.shapes):
                if getattr(form, "has_text_frame", False) and form.text_frame.text.strip():
                    teile.append(f"  [Element {index}] " + form.text_frame.text.strip().replace("\n", " / "))
                elif getattr(form, "has_table", False) and form.has_table:
                    teile.append(f"  [Element {index}, Tabelle] " + " | ".join(z.text for r in form.table.rows for z in r.cells)[:400])
                elif getattr(form, "shape_type", None) == 13:
                    teile.append(f"  [Element {index}, Bild]")
            notizen = ""
            try:
                if folie.has_notes_slide:
                    notizen = folie.notes_slide.notes_text_frame.text.strip()
            except Exception:
                notizen = ""
            zeilen.append(f"Folie {nummer}:\n" + "\n".join(teile) + (f"\n  Notizen: {notizen}" if notizen else ""))
        inhalt = f"PowerPoint mit {len(praesentation.slides)} Folien. Bearbeiten mit edit_pptx (Folien ab 1, Element-Nummern wie hier angegeben).\n\n" + "\n\n".join(zeilen)
        return {"kind": "praesentation", "name": name, "slides": len(praesentation.slides), "content": inhalt[:MAX_TEXT]}

    def _extract_office(self, name: str, pfad: Path, art: str) -> dict:
        endung = _endung(name)
        if endung in ("odt", "ods"):
            return self._extract_odf(name, pfad, art)
        if endung == "rtf":
            roh = pfad.read_bytes()[:400_000].decode("latin-1", errors="replace")
            import re

            text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", roh)
            return {"kind": art, "name": name, "content": text[:MAX_TEXT]}
        from app.services.dateiansicht_service import ansehen

        daten = ansehen(pfad)
        if daten.get("art") == "fehler":
            return {"error": daten.get("fehler", "unbekannt")}
        if daten.get("art") == "tabelle":
            teile = []
            for blatt in daten.get("blaetter", []):
                zeilen = [" | ".join(z) for z in blatt.get("zeilen", [])]
                teile.append(f"Blatt „{blatt.get('name', '')}“:\n" + "\n".join(zeilen))
            return {"kind": "tabelle", "name": name, "content": "\n\n".join(teile)[:MAX_TEXT]}
        text = str(daten.get("text", ""))
        if endung == "docx":
            text = "Word-Dokument. Bearbeiten mit edit_docx.\n\n" + text
        return {"kind": "dokument", "name": name, "content": text[:MAX_TEXT]}

    def _extract_odf(self, name: str, pfad: Path, art: str) -> dict:
        try:
            from odf import teletype, text
            from odf.opendocument import load

            dokument = load(str(pfad))
            teile = [teletype.extractText(e).strip() for e in dokument.getElementsByType(text.P) + dokument.getElementsByType(text.H)]
        except Exception as exc:
            return {"error": str(exc)}
        return {"kind": art, "name": name, "content": "\n".join(t for t in teile if t)[:MAX_TEXT]}

    def _extract_text(self, name: str, pfad: Path) -> dict:
        with pfad.open("rb") as datei:
            roh = datei.read(500_000)
        if b"\x00" in roh[:4000]:
            return {"kind": "datei", "name": name, "content": f"Binärdatei „{name}“ ({pfad.stat().st_size} Bytes). Inhalt nicht als Text lesbar."}
        text = roh.decode("utf-8", errors="replace")
        return {"kind": "text", "name": name, "content": text[:MAX_TEXT]}

    async def beschreiben(self, pfad: Path, mime: str = "", provider_name: str | None = None) -> dict:
        ergebnis = await self._describe_image(pfad.name, mime or f"image/{_endung(pfad.name).replace('jpg', 'jpeg')}", pfad, provider_name)
        ergebnis["pfad"] = str(pfad)
        return ergebnis

    async def _describe_image(
        self, name: str, mime: str, pfad: Path, provider_name: str | None
    ) -> dict:
        settings = get_settings()
        user = get_settings_service()
        saved_provider, saved_model = user.selection()
        provider_name = provider_name or saved_provider or settings.default_provider
        vision_model = (
            user.get().get("vision_model")
            or VISION_DEFAULTS.get(provider_name)
            or saved_model
            or settings.jon_model
        )
        from app.services.aufgaben_modelle import sehen_wahl

        provider_name, vision_model = sehen_wahl(provider_name, vision_model)
        provider = get_registry().all().get(provider_name)
        hinweis = _bildhinweis(name, pfad)
        if not isinstance(provider, OpenAICompatibleProvider) or not provider.available():
            return hinweis
        try:
            daten, art = await asyncio.to_thread(_verkleinern, pfad, mime)
        except Exception as fehler:
            leise(fehler, "services/attachment_service")
            return hinweis
        data_url = f"data:{art};base64,{base64.b64encode(daten).decode('ascii')}"
        try:
            text = await asyncio.wait_for(provider.describe_image(vision_model, data_url, IMAGE_PROMPT, max_tokens=700), 60)
        except Exception as exc:
            leise(exc, "services/attachment_service")
            return hinweis
        return {"kind": "image", "name": name, "content": text.strip()[:MAX_TEXT]}


def _bildhinweis(name: str, pfad: Path) -> dict:
    return {"kind": "image", "name": name, "content": f"Bild „{name}“ wurde hochgeladen. Mit look_at_image (Pfad {pfad}) kannst du es dir ansehen."}


def _verkleinern(pfad: Path, mime: str) -> tuple[bytes, str]:
    from PIL import Image

    with Image.open(pfad) as bild:
        if max(bild.size) <= BILD_KANTE and pfad.stat().st_size < 1_500_000 and (mime or "").startswith("image/") and mime != "image/bmp":
            return pfad.read_bytes(), mime
        bild.thumbnail((BILD_KANTE, BILD_KANTE))
        ausgabe = io.BytesIO()
        bild.convert("RGB").save(ausgabe, format="JPEG", quality=85)
        return ausgabe.getvalue(), "image/jpeg"


_service: AttachmentService | None = None


def get_attachment_service() -> AttachmentService:
    global _service
    if _service is None:
        _service = AttachmentService()
    return _service
