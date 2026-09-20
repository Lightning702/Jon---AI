from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

from app.services import verknuepfung_service as vk
from app.services.browserwahl import JON, SYSTEM, aufloesen
from app.services.pptx_effekte import EINFACH, MODERN, animieren, uebergang
from app.services.pptx_inhalte import punkt_text, zeilen_aus
from app.services.pptx_service import get_pptx_service


@pytest.mark.parametrize(
    "eingabe,erwartet",
    [
        ("edge", "edge"),
        ("Microsoft Edge", "edge"),
        ("nutze brave", "brave"),
        ("mit Brave Browser", "brave"),
        ("Google Chrome", "chrome"),
        ("firefox", "firefox"),
        ("standardbrowser", SYSTEM),
        ("nimm meinen normalen Browser", SYSTEM),
        ("in meinem gewohnten Browser", SYSTEM),
        ("jon", JON),
        ("eigener", JON),
        ("", ""),
        ("irgendwas", ""),
    ],
)
def test_browsernamen_aus_normaler_sprache(eingabe, erwartet):
    assert aufloesen(eingabe) == erwartet


def test_jons_browser_ist_der_standard(monkeypatch):
    from app.services import browserwahl
    from app.services.settings_service import DEFAULTS

    assert DEFAULTS["web_browser"] == JON

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    monkeypatch.setattr(
        browserwahl.webbrowser,
        "open",
        lambda url: pytest.fail("der Systembrowser haette nicht laufen duerfen"),
    )
    aufgerufen: dict = {}

    def _goto(aktion, args):
        aufgerufen["aktion"] = aktion
        aufgerufen["url"] = args.get("url")
        return {"ok": True}

    monkeypatch.setattr("app.services.browser.werkzeuge.ausfuehren", _goto)
    ergebnis = browserwahl.oeffnen("example.com")
    assert aufgerufen["aktion"] == "goto"
    assert aufgerufen["url"] == "https://example.com"
    assert ergebnis["browser"] == "Jons privater Browser"


def test_normaler_browser_auf_ausdruecklichen_wunsch(monkeypatch):
    from app.services import browserwahl

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    monkeypatch.setattr(
        "app.services.browser.werkzeuge.ausfuehren",
        lambda aktion, args: pytest.fail("Jons Browser haette nicht laufen duerfen"),
    )
    geoeffnet: dict = {}
    monkeypatch.setattr(
        browserwahl.webbrowser,
        "open",
        lambda url: geoeffnet.setdefault("url", url) or True,
    )
    ergebnis = browserwahl.oeffnen("example.com", "nimm meinen normalen Browser")
    assert geoeffnet["url"] == "https://example.com"
    assert ergebnis["browser"] == "Standardbrowser"


def test_ausdruecklicher_wunsch_schlaegt_den_standard(monkeypatch):
    from app.services import browserwahl

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    monkeypatch.setattr(browserwahl, "_pfad", lambda schluessel: "")
    geoeffnet: dict = {}
    monkeypatch.setattr(
        browserwahl.webbrowser, "open", lambda url: geoeffnet.setdefault("url", url) or True
    )
    ergebnis = browserwahl.oeffnen("example.com", "nutze Edge")
    assert geoeffnet["url"] == "https://example.com"
    assert "Microsoft Edge" in ergebnis.get("hinweis", "")


def test_web_search_ohne_wunsch_bleibt_direkt(monkeypatch):
    import asyncio

    from app.services import browserwahl
    from app.services.tools import ToolBox

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)

    def _verboten(frage, anzahl):
        raise AssertionError("Ohne Auftrag darf kein Browser starten.")

    async def _direkt(frage, anzahl=6, read=False):
        return {"treffer": [{"title": "Direkt"}]}

    monkeypatch.setattr("app.services.websuche_browser.suchen", _verboten)
    monkeypatch.setattr("app.services.websearch_service.search_web", _direkt)
    daten = json.loads(asyncio.run(ToolBox().execute("web_search", {"query": "pizza"})))
    assert daten["browser"] == "Direktsuche"
    assert daten["treffer"]
    assert "hinweis" not in daten


def test_web_search_nimmt_jons_browser_auf_ansage(monkeypatch):
    import asyncio

    from app.services import browserwahl
    from app.services.tools import ToolBox

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    monkeypatch.setattr(
        "app.services.websuche_browser.suchen",
        lambda frage, anzahl: {"treffer": [{"title": "Treffer", "url": "x"}]},
    )
    roh = asyncio.run(
        ToolBox().execute("web_search", {"query": "pizza", "browser": "jon"})
    )
    daten = json.loads(roh)
    assert daten["browser"] == "Jons privater Browser"
    assert daten["treffer"]


def test_web_search_sagt_wenn_der_rueckfall_greift(monkeypatch):
    import asyncio

    from app.services import browserwahl
    from app.services.tools import ToolBox

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)

    def _kaputt(frage, anzahl):
        raise RuntimeError("Captcha")

    async def _direkt(frage, anzahl=6, read=False):
        return {"treffer": [{"title": "Direkt"}, {"title": "Zweiter"}], "mager": False}

    monkeypatch.setattr("app.services.websuche_browser.suchen", _kaputt)
    monkeypatch.setattr("app.services.websearch_service.search_web", _direkt)
    daten = json.loads(
        asyncio.run(ToolBox().execute("web_search", {"query": "x", "browser": "jon"}))
    )
    assert daten["browser"] == "Direktsuche"
    assert daten["treffer"][0]["title"] == "Direkt"


def test_web_search_mit_fremdem_browser_oeffnet_dort(monkeypatch):
    import asyncio

    from app.services import browserwahl
    from app.services.tools import ToolBox

    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    monkeypatch.setattr(
        browserwahl, "oeffnen",
        lambda url, erzwinge="": {"ok": True, "geoeffnet": url, "browser": "Brave"},
    )
    daten = json.loads(
        asyncio.run(ToolBox().execute("web_search", {"query": "pizza", "browser": "brave"}))
    )
    assert "duckduckgo" in daten["geoeffnet"]
    assert "nicht mitlesen" in daten["hinweis"]


def test_open_url_kennt_den_browser_parameter():
    from app.services.tools import ToolBox

    schema = {t["function"]["name"]: t for t in ToolBox()._eigene_tools()}
    felder = schema["open_url"]["function"]["parameters"]["properties"]
    assert "browser" in felder
    assert "JONS PRIVATEM" in schema["open_url"]["function"]["description"]
    assert "browser" in schema["web_search"]["function"]["parameters"]["properties"]


def test_systemprompt_haelt_die_browserregel_fest():
    from app.services.chat_service import SYSTEM_PROMPT

    assert "JONS PRIVATER BROWSER" in SYSTEM_PROMPT
    assert "AUSNAHME" in SYSTEM_PROMPT
    assert "KEIN FENSTER OHNE AUFTRAG" in SYSTEM_PROMPT


def test_reine_recherche_bekommt_keine_browserwerkzeuge():
    from app.services.tools import ToolBox

    box = ToolBox()
    for frage in (
        "was kostet die rtx 5090 aktuell",
        "such mir infos ueber quantencomputer",
        "wann ist die naechste sonnenfinsternis",
    ):
        namen = [t["function"]["name"] for t in box.schema(frage)]
        assert "web_search" in namen
        assert not [n for n in namen if n.startswith("browser_")]
    for auftrag in ("oeffne mir youtube im browser", "klick auf den knopf"):
        namen = [t["function"]["name"] for t in box.schema(auftrag)]
        assert [n for n in namen if n.startswith("browser_")]


@pytest.fixture()
def folien():
    return [
        {
            "layout": "title",
            "title": "Probe",
            "subtitle": "Untertitel",
            "notes": "Einstieg.",
        },
        {
            "layout": "agenda",
            "title": "Ablauf",
            "text": "Drei Teile.",
            "items": [
                {"titel": "Erstens", "text": "Erklaerung dazu in einem ganzen Satz."},
                {"titel": "Zweitens", "text": "Noch eine Erklaerung mit Inhalt."},
            ],
        },
        {
            "layout": "bullets",
            "title": "Punkte",
            "text": "Eine Einleitung, die den Gedanken wirklich erklaert und dabei "
            "ueber eine Zeile hinausgeht, damit der Umbruch geprueft wird.",
            "bullets": [
                "Erster Punkt: eine ausformulierte Aussage, die deutlich laenger ist "
                "als ein Stichwort und damit umbricht",
                "Zweiter Punkt: ebenfalls ein vollstaendiger Satz mit Inhalt",
                "Dritter Punkt: und noch einer, damit die Karte wachsen muss",
            ],
        },
        {
            "layout": "chart",
            "title": "Zahlen",
            "text": "Der Verlauf.",
            "diagramm": {
                "art": "balken",
                "kategorien": ["2022", "2023", "2024"],
                "reihen": [{"name": "Wert", "werte": [10, 20, 30]}],
            },
        },
        {
            "layout": "table",
            "title": "Tabelle",
            "tabelle": [["A", "B"], ["1", "2"], ["3", "4"]],
        },
        {
            "layout": "text",
            "title": "Fliesstext",
            "absaetze": ["Erster Absatz mit Inhalt.", "Zweiter Absatz mit Inhalt."],
        },
        {
            "layout": "compare",
            "title": "Vergleich",
            "items": [
                {"titel": "Dafuer", "text": "Gruende", "bullets": ["Punkt eins"]},
                {"titel": "Dagegen", "text": "Gegengruende", "bullets": ["Punkt zwei"]},
            ],
        },
    ]


def test_praesentation_bekommt_uebergaenge_und_animationen(folien, tmp_path):
    ziel = tmp_path / "probe.pptx"
    ergebnis = get_pptx_service().create("Probe", folien, str(ziel), "forest")
    assert not ergebnis.get("error"), ergebnis
    assert ergebnis["uebergaenge"] == ergebnis["slides"]
    assert ergebnis["animationen"] > 0
    with zipfile.ZipFile(ziel) as z:
        namen = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        assert len(namen) == ergebnis["slides"]
        for name in namen:
            xml = z.read(name).decode("utf-8")
            ElementTree.fromstring(xml)
            assert "<p:transition" in xml
            assert "<p:timing" in xml


def test_effekte_lassen_sich_abschalten(folien, tmp_path):
    ziel = tmp_path / "ohne.pptx"
    ergebnis = get_pptx_service().create(
        "Ohne", folien, str(ziel), "ocean", effekte=False
    )
    assert ergebnis["uebergaenge"] == 0
    assert ergebnis["animationen"] == 0
    with zipfile.ZipFile(ziel) as z:
        xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "<p:transition" not in xml


def test_diagramm_und_tabelle_sind_echt(folien, tmp_path):
    from pptx import Presentation

    ziel = tmp_path / "inhalt.pptx"
    get_pptx_service().create("Inhalt", folien, str(ziel), "ocean")
    prs = Presentation(str(ziel))
    diagramme = [f for s in prs.slides for f in s.shapes if f.has_chart]
    tabellen = [f for s in prs.slides for f in s.shapes if f.has_table]
    assert len(diagramme) == 1
    assert len(tabellen) == 1
    werte = list(diagramme[0].chart.plots[0].series[0].values)
    assert werte == [10.0, 20.0, 30.0]
    assert tabellen[0].table.cell(0, 0).text == "A"


def test_alle_neuen_layouts_bauen_durch(folien, tmp_path):
    ergebnis = get_pptx_service().create(
        "Alle", folien, str(tmp_path / "alle.pptx"), "teal"
    )
    assert ergebnis["layouts"] == [
        "title",
        "agenda",
        "bullets",
        "chart",
        "table",
        "text",
        "compare",
    ]


def test_kaputtes_diagramm_faellt_auf_punkte_zurueck(tmp_path):
    from pptx import Presentation

    folie = [
        {
            "layout": "chart",
            "title": "Leer",
            "diagramm": {"art": "balken", "kategorien": [], "reihen": []},
            "bullets": ["Ersatzpunkt mit Inhalt"],
        }
    ]
    ziel = tmp_path / "ersatz.pptx"
    get_pptx_service().create("Ersatz", folie, str(ziel), "coral")
    prs = Presentation(str(ziel))
    assert not [f for s in prs.slides for f in s.shapes if f.has_chart]
    text = " ".join(
        f.text_frame.text for s in prs.slides for f in s.shapes
        if getattr(f, "has_text_frame", False)
    )
    assert "Ersatzpunkt" in text


def test_titel_taucht_nicht_doppelt_auf(tmp_path):
    from pptx import Presentation

    folie = [{"layout": "bullets", "title": "Thema", "bullets": ["Ein Punkt"]}]
    ziel = tmp_path / "titel.pptx"
    get_pptx_service().create("Thema", folie, str(ziel), "gold")
    prs = Presentation(str(ziel))
    zweite = list(prs.slides)[1]
    text = " ".join(
        f.text_frame.text for f in zweite.shapes if getattr(f, "has_text_frame", False)
    )
    assert text.count("Thema") == 1


def test_lange_punkte_bekommen_kleinere_schrift():
    from app.services.pptx_service import _bullet_groesse

    kurz = ["Kurz", "Auch kurz"]
    lang = ["Ein sehr langer Punkt " * 12 for _ in range(6)]
    assert _bullet_groesse(kurz, 11.9, 4.0) >= _bullet_groesse(lang, 11.9, 4.0)
    assert _bullet_groesse(lang, 11.9, 4.0) >= 13


def test_zeilen_werden_mit_umbruch_geschaetzt():
    from app.services.pptx_service import _zeilen

    assert _zeilen("kurz", 11.9, 19) == 1
    assert _zeilen("x" * 400, 11.9, 19) > 1


@pytest.mark.parametrize("art", sorted(set(list(EINFACH) + list(MODERN))))
def test_jeder_uebergang_erzeugt_gueltiges_xml(art, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    assert uebergang(folie, art) is True
    ziel = tmp_path / f"{art}.pptx"
    prs.save(str(ziel))
    with zipfile.ZipFile(ziel) as z:
        ElementTree.fromstring(z.read("ppt/slides/slide1.xml"))


def test_unbekannter_uebergang_wird_abgelehnt(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    assert uebergang(folie, "gibtsnicht") is False
    assert uebergang(folie, "keiner") is True


def test_animation_ohne_formen_bleibt_leer(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    assert animieren(folie, []) == 0


def test_uebergang_steht_vor_dem_timing(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    kasten = folie.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    kasten.text_frame.text = "Text"
    animieren(folie, [{"form": kasten}])
    uebergang(folie, "fade")
    ziel = tmp_path / "reihenfolge.pptx"
    prs.save(str(ziel))
    with zipfile.ZipFile(ziel) as z:
        xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
    assert xml.index("<p:transition") < xml.index("<p:timing")


def test_zeilen_aus_versteht_verschiedene_formen():
    assert zeilen_aus([["a", "b"], ["c", "d"]]) == [["a", "b"], ["c", "d"]]
    assert zeilen_aus([{"x": 1, "y": 2}]) == [["x", "y"], ["1", "2"]]
    assert zeilen_aus({"Jahr": ["2023", "2024"], "Wert": ["1", "2"]}) == [
        ["Jahr", "Wert"],
        ["2023", "1"],
        ["2024", "2"],
    ]


def test_punkt_text_trennt_begriff_und_erklaerung():
    assert punkt_text("Begriff: Erklaerung") == ("Begriff", "Erklaerung")
    assert punkt_text({"titel": "T", "text": "E"}) == ("T", "E")
    assert punkt_text("nur text") == ("nur text", "")


def test_skill_verlangt_ausformulierten_text():
    inhalt = Path("skills/powerpoint.md")
    if not inhalt.is_file():
        inhalt = Path(__file__).resolve().parents[2] / "skills" / "powerpoint.md"
    text = inhalt.read_text(encoding="utf-8")
    assert "Niemals Platzhalter" in text
    assert "bild_prompt" in text
    assert "diagramm" in text
    assert "400–900 Zeichen" in text


def test_verknuepfung_wird_angelegt_und_entfernt(tmp_path, monkeypatch):
    monkeypatch.setattr(vk, "schreibtisch", lambda: tmp_path)
    programm = tmp_path / "Jon.exe"
    programm.write_bytes(b"MZ")
    assert vk.vorhanden() == ""
    ergebnis = vk.anlegen(str(programm))
    if ergebnis.get("error") and sys.platform != "win32":
        pytest.skip("Verknuepfungen brauchen eine Desktop-Umgebung")
    assert ergebnis.get("ok") is True, ergebnis
    assert Path(ergebnis["pfad"]).exists()
    assert vk.vorhanden() == ergebnis["pfad"]
    assert vk.entfernen()["entfernt"] is True
    assert vk.vorhanden() == ""


def test_verknuepfung_haelt_einen_ordner_nicht_fuer_sich_selbst(tmp_path, monkeypatch):
    monkeypatch.setattr(vk, "schreibtisch", lambda: tmp_path)
    (tmp_path / "Jon").mkdir()
    assert vk.vorhanden() == ""


def test_verknuepfung_ohne_programm_meldet_sich_verstaendlich(tmp_path, monkeypatch):
    monkeypatch.setattr(vk, "schreibtisch", lambda: tmp_path)
    monkeypatch.setattr(vk, "_exe", lambda: None)
    ergebnis = vk.anlegen()
    assert "Jon.exe" in ergebnis["error"]


def test_verknuepfung_werkzeug_ist_angemeldet():
    from app.services.tools import describe_tool, werkzeugnamen
    from app.services.werkzeug_register import finden

    assert "desktop_verknuepfung" in werkzeugnamen()
    assert finden("desktop_verknuepfung")
    assert "Desktop" in describe_tool("desktop_verknuepfung", {})
    assert "Desktop" in describe_tool("desktop_verknuepfung", {"aktion": "entfernen"})


def test_zip_bekommt_den_einrichter():
    wurzel = Path(__file__).resolve().parents[2]
    text = (wurzel / "scripts" / "build_installer.py").read_text(encoding="utf-8")
    assert "Jon einrichten.exe" in text
    assert "build_einrichter" in text
    assert "--icon" in text
    assert "LIESMICH.txt" in text
    assert (wurzel / "frontend" / "electron" / "icon.ico").is_file()


def test_einrichter_findet_seinen_ordner_und_meldet_fehlendes_jon(tmp_path, monkeypatch):
    import importlib.util as u

    wurzel = Path(__file__).resolve().parents[2]
    quelle = wurzel / "scripts" / "jon_einrichten.py"
    assert quelle.is_file()
    spec = u.spec_from_file_location("jon_einrichten", quelle)
    modul = u.module_from_spec(spec)
    spec.loader.exec_module(modul)

    monkeypatch.setattr(modul, "_ordner", lambda: tmp_path)
    monkeypatch.setattr(modul, "schreibtisch", lambda: tmp_path / "desktop")
    monkeypatch.setattr(modul, "startmenue", lambda: tmp_path / "menue")
    monkeypatch.setattr(modul, "_warten", lambda text="": None)
    ausgabe: list[str] = []
    monkeypatch.setattr(modul, "_sagen", lambda text="": ausgabe.append(text))
    assert modul.main() == 1
    assert any("Jon.exe liegt nicht neben" in z for z in ausgabe)
    assert any(str(tmp_path) in z for z in ausgabe)


def test_einrichter_legt_beide_verknuepfungen_an(tmp_path, monkeypatch):
    import importlib.util as u

    if sys.platform != "win32":
        pytest.skip("Verknuepfungen sind hier Windows-Sache")
    wurzel = Path(__file__).resolve().parents[2]
    spec = u.spec_from_file_location(
        "jon_einrichten", wurzel / "scripts" / "jon_einrichten.py"
    )
    modul = u.module_from_spec(spec)
    spec.loader.exec_module(modul)

    (tmp_path / "Jon.exe").write_bytes(b"MZ")
    schreibtisch = tmp_path / "desktop"
    menue = tmp_path / "menue"
    schreibtisch.mkdir()
    menue.mkdir()
    gestartet: list[Path] = []
    monkeypatch.setattr(modul, "_ordner", lambda: tmp_path)
    monkeypatch.setattr(modul, "schreibtisch", lambda: schreibtisch)
    monkeypatch.setattr(modul, "startmenue", lambda: menue)
    monkeypatch.setattr(modul, "_sagen", lambda text="": None)
    monkeypatch.setattr(modul, "_warten", lambda text="": None)
    monkeypatch.setattr(
        modul, "starten", lambda exe, ordner: gestartet.append(exe) or ""
    )
    monkeypatch.setattr(modul.time, "sleep", lambda sekunden: None)
    assert modul.main() == 0
    assert (schreibtisch / "Jon.lnk").is_file()
    assert (menue / "Jon.lnk").is_file()
    assert gestartet and gestartet[0].name == "Jon.exe"

    ziel = schreibtisch / "Jon.lnk"
    assert ziel.stat().st_size > 200


def test_app_bietet_die_verknuepfung_selbst_an():
    wurzel = Path(__file__).resolve().parents[2]
    text = (wurzel / "frontend" / "electron" / "main.cjs").read_text(encoding="utf-8")
    assert "verknuepfungAnbieten" in text
    assert "writeShortcutLink" in text
    assert "portabelInstalliert" in text
    assert "Nicht mehr fragen" in text


def test_electron_oeffnet_links_ueber_jons_browser():
    wurzel = Path(__file__).resolve().parents[2]
    text = (wurzel / "frontend" / "electron" / "main.cjs").read_text(encoding="utf-8")
    assert "adresseOeffnen" in text
    assert "/browser/oeffnen" in text
    assert "istEigeneSeite" in text
    stelle = text.index("mainWindow.webContents.setWindowOpenHandler")
    ausschnitt = text[stelle : stelle + 200]
    assert "adresseOeffnen" in ausschnitt
    assert "shell.openExternal" not in ausschnitt


def test_route_oeffnet_ueber_die_browserwahl(monkeypatch):
    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.services import browserwahl

    gerufen: dict = {}

    def _oeffnen(url, erzwinge=""):
        gerufen["url"] = url
        gerufen["erzwinge"] = erzwinge
        return {"ok": True, "geoeffnet": url, "browser": "Jons privater Browser"}

    monkeypatch.setattr(browserwahl, "oeffnen", _oeffnen)
    with TestClient(create_app()) as client:
        antwort = client.post("/api/browser/oeffnen", json={"url": "example.com"})
    assert antwort.status_code == 200
    assert antwort.json()["browser"] == "Jons privater Browser"
    assert gerufen["url"] == "example.com"
    assert gerufen["erzwinge"] == ""


def test_route_lehnt_leere_adresse_ab():
    from fastapi.testclient import TestClient

    from app.main import create_app

    with TestClient(create_app()) as client:
        assert client.post("/api/browser/oeffnen", json={"url": " "}).status_code == 400


def test_oberflaeche_loest_namen_auf():
    from app.services.oberflaeche_service import aufloesen

    for eingabe, erwartet in (
        ("tresor", "tresor"),
        ("/vault", "tresor"),
        ("passwort", "tresor"),
        ("mach die karten auf", "maps"),
        ("ziele", "denken"),
        ("aufgabenliste", "aufgaben"),
        ("quiz", "lernen"),
        ("", ""),
        ("voelliger unsinn", ""),
    ):
        assert aufloesen(eingabe) == erwartet, eingabe


def test_oberflaeche_werkzeug_meldet_das_ziel():
    import asyncio
    import json as _json

    from app.services.tools import ToolBox

    daten = _json.loads(
        asyncio.run(ToolBox().execute("oberflaeche", {"werkzeug": "tresor"}))
    )
    assert daten["ok"] is True
    assert daten["oeffne"] == "tresor"
    assert daten["befehl"] == "/tresor"
    assert "Passwort-Tresor" in daten["name"]

    liste = _json.loads(asyncio.run(ToolBox().execute("oberflaeche", {})))
    assert len(liste["werkzeuge"]) > 20

    unbekannt = _json.loads(
        asyncio.run(ToolBox().execute("oberflaeche", {"werkzeug": "raumschiff"}))
    )
    assert unbekannt.get("error")
    assert "moeglich" in unbekannt


def test_jede_oberflaeche_hat_eine_aktion_im_frontend():
    from app.services.oberflaeche_service import ZIELE

    wurzel = Path(__file__).resolve().parents[2]
    text = (wurzel / "frontend" / "src" / "App.tsx").read_text(encoding="utf-8")
    stelle = text.index("const oberflaecheOeffnen")
    block = text[stelle : text.index("aktionen[ziel]", stelle)]
    for ziel in ZIELE:
        assert f"{ziel}:" in block, ziel


def test_chat_reicht_das_oeffnen_durch():
    wurzel = Path(__file__).resolve().parents[2]
    text = (
        wurzel / "backend" / "app" / "services" / "chat_service.py"
    ).read_text(encoding="utf-8")
    assert 'chunk.name == "oberflaeche"' in text
    assert 'event["oeffne"]' in text
    api = (wurzel / "frontend" / "src" / "lib" / "api.ts").read_text(encoding="utf-8")
    assert "oeffne?: string;" in api


@pytest.mark.parametrize(
    "werkzeug,args,erwartet",
    [
        ("start_program", {"path": "https://www.youtube.com"}, ("https://www.youtube.com", "")),
        ("start_program", {"path": "youtube.com"}, ("youtube.com", "")),
        ("start_program", {"path": "www.youtube.com"}, ("www.youtube.com", "")),
        (
            "start_program",
            {"path": "chrome", "args": ["https://youtube.com"]},
            ("https://youtube.com", "chrome"),
        ),
        (
            "start_program",
            {"path": "msedge.exe", "args": ["youtube.com"]},
            ("youtube.com", "edge"),
        ),
        (
            "run_powershell",
            {"command": 'Start-Process "https://www.youtube.com"'},
            ("https://www.youtube.com", ""),
        ),
        (
            "run_powershell",
            {"command": 'Start-Process -FilePath "https://youtube.com"'},
            ("https://youtube.com", ""),
        ),
        (
            "run_powershell",
            {"command": "Start-Process msedge https://youtube.com"},
            ("https://youtube.com", "edge"),
        ),
        ("run_cmd", {"command": "start https://www.youtube.com"}, ("https://www.youtube.com", "")),
        ("run_cmd", {"command": 'start "" "https://youtube.com"'}, ("https://youtube.com", "")),
        ("run_cmd", {"command": "explorer https://youtube.com"}, ("https://youtube.com", "")),
        ("run_cmd", {"command": "start youtube.com"}, ("youtube.com", "")),
    ],
)
def test_webadressen_werden_umgeleitet(werkzeug, args, erwartet):
    from app.services.weboeffnen import pruefen

    assert pruefen(werkzeug, args) == erwartet


@pytest.mark.parametrize(
    "werkzeug,args",
    [
        ("start_program", {"path": "notepad.exe"}),
        ("start_program", {"path": "C:/Programme/Jon/Jon.exe"}),
        ("start_program", {"path": "command.com"}),
        ("run_powershell", {"command": "Get-Process | Where-Object Name -eq chrome"}),
        ("run_powershell", {"command": "Start-Process notepad.exe"}),
        ("run_powershell", {"command": 'Start-Process "C:/Temp/bericht.pdf"'}),
        ("run_cmd", {"command": "dir C:/"}),
        ("run_cmd", {"command": "start notepad && dir"}),
        ("write_file", {"path": "https://youtube.com"}),
        ("run_powershell", {"command": "Start-Process winword.exe"}),
    ],
)
def test_echte_shell_befehle_bleiben_shell(werkzeug, args):
    from app.services.weboeffnen import pruefen

    assert pruefen(werkzeug, args) is None


def test_eigene_dateien_gelten_nicht_als_adresse(tmp_path, monkeypatch):
    from app.services import weboeffnen

    datei = tmp_path / "bericht.com"
    datei.write_text("x", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    assert weboeffnen.ist_adresse("bericht.com") is False
    assert weboeffnen.ist_adresse("youtube.com") is True


def test_umleitung_laeuft_ueber_die_browserwahl(monkeypatch):
    from app.services import browserwahl, weboeffnen

    gerufen: dict = {}

    def _oeffnen(url, erzwinge=""):
        gerufen["url"] = url
        gerufen["erzwinge"] = erzwinge
        return {"ok": True, "geoeffnet": url, "browser": "Jons privater Browser"}

    monkeypatch.setattr(browserwahl, "oeffnen", _oeffnen)
    monkeypatch.setattr(browserwahl, "wahl", lambda: JON)
    ergebnis = weboeffnen.umleiten(
        "run_powershell", {"command": "Start-Process https://youtube.com"}
    )
    assert ergebnis["ok"] is True
    assert gerufen["url"] == "https://youtube.com"
    assert ergebnis["umgeleitet_von"] == "run_powershell"
    assert "Jons privater Browser" in ergebnis["hinweis"]


def test_execute_leitet_um_statt_die_shell_zu_starten(monkeypatch):
    import asyncio
    import json as _json

    from app.services import weboeffnen
    from app.services.tools import ToolBox

    monkeypatch.setattr(
        weboeffnen,
        "umleiten",
        lambda name, args: {"ok": True, "geoeffnet": "x", "browser": "Jons privater Browser"},
    )

    def _nie(self, name, args):
        raise AssertionError("die Shell haette nicht laufen duerfen")

    monkeypatch.setattr("app.services.tools.ToolBox._dispatch", _nie)
    daten = _json.loads(
        asyncio.run(
            ToolBox().execute("run_cmd", {"command": "start https://youtube.com"})
        )
    )
    assert daten["browser"] == "Jons privater Browser"


def test_systemprompt_verbietet_die_shell_fuer_webseiten():
    from app.services.chat_service import SYSTEM_PROMPT

    assert "IMMER open_url" in SYSTEM_PROMPT
    assert "niemals start_program" in SYSTEM_PROMPT
