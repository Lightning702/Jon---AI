from pathlib import Path

import pytest

from app.services import dateiansicht_service as ansicht


def test_markdown_und_code(tmp_path):
    md = tmp_path / "notiz.md"
    md.write_text("# Titel\n\nText", encoding="utf-8")
    py = tmp_path / "skript.py"
    py.write_text("print('hi')", encoding="utf-8")
    assert ansicht.ansehen(md)["art"] == "markdown"
    code = ansicht.ansehen(py)
    assert code["art"] == "code" and code["sprache"] == "py" and "print" in code["text"]


def test_csv_wird_tabelle(tmp_path):
    datei = tmp_path / "werte.csv"
    datei.write_text("Name;Wert\nA;1\nB;2\n", encoding="utf-8")
    ergebnis = ansicht.ansehen(datei)
    assert ergebnis["art"] == "tabelle"
    assert ergebnis["blaetter"][0]["zeilen"][:2] == [["Name", "Wert"], ["A", "1"]]


def test_word_wird_markdown(tmp_path):
    docx = pytest.importorskip("docx")
    dokument = docx.Document()
    dokument.add_heading("Bericht", level=1)
    dokument.add_paragraph("Erster Punkt", style="List Bullet")
    tabelle = dokument.add_table(rows=2, cols=2)
    tabelle.cell(0, 0).text = "Kopf"
    tabelle.cell(1, 1).text = "Wert"
    ziel = tmp_path / "bericht.docx"
    dokument.save(ziel)
    text = ansicht.ansehen(ziel)["text"]
    assert "## Bericht" in text
    assert "- Erster Punkt" in text
    assert "| Kopf |" in text


def test_excel_wird_tabelle(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    blatt.title = "Umsatz"
    blatt.append(["Monat", "Betrag"])
    blatt.append(["Jan", 12.5])
    ziel = tmp_path / "umsatz.xlsx"
    mappe.save(ziel)
    ergebnis = ansicht.ansehen(ziel)
    assert ergebnis["art"] == "tabelle"
    assert ergebnis["blaetter"][0]["name"] == "Umsatz"
    assert ergebnis["blaetter"][0]["zeilen"][1] == ["Jan", "12.5"]


def test_praesentation_mit_formen_und_bild(tmp_path):
    pptx = pytest.importorskip("pptx")
    from PIL import Image
    from pptx.util import Inches, Pt

    bild = tmp_path / "bild.png"
    Image.new("RGB", (40, 30), "red").save(bild)
    praesentation = pptx.Presentation()
    folie = praesentation.slides.add_slide(praesentation.slide_layouts[6])
    feld = folie.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    feld.text_frame.text = "Hallo Jon"
    feld.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    folie.shapes.add_picture(str(bild), Inches(5), Inches(2))
    folie.notes_slide.notes_text_frame.text = "Sprechernotiz"
    ziel = tmp_path / "deck.pptx"
    praesentation.save(ziel)
    ergebnis = ansicht.ansehen(ziel)
    assert ergebnis["art"] == "folien"
    formen = ergebnis["folien"][0]["formen"]
    text = next(f for f in formen if f.get("absaetze"))
    assert text["absaetze"][0]["text"] == "Hallo Jon"
    assert text["absaetze"][0]["groesse"] == 32
    bildform = next(f for f in formen if f["art"] == "bild")
    daten, mime = ansicht.folienbild(ziel, 1, bildform["index"])
    assert mime in ("image/png", "image/jpeg") and len(daten) > 50
    assert ergebnis["folien"][0]["notizen"] == "Sprechernotiz"


def test_pdf_ohne_seitenbilder(tmp_path, monkeypatch):
    pypdf = pytest.importorskip("pypdf")
    schreiber = pypdf.PdfWriter()
    schreiber.add_blank_page(width=200, height=200)
    ziel = tmp_path / "leer.pdf"
    with ziel.open("wb") as datei:
        schreiber.write(datei)
    monkeypatch.setattr(ansicht.shutil, "which", lambda name: None)
    ergebnis = ansicht.ansehen(ziel)
    assert ergebnis["art"] == "pdf"
    assert ergebnis["seiten"] == 1 and ergebnis["bilder"] is False


def test_relativer_pfad_sucht_zuerst_in_jons_ordner(tmp_path, monkeypatch):
    from app.api import mobile_routes

    wurzel = tmp_path / "Jon"
    wurzel.mkdir()
    (wurzel / "Notiz.md").write_text("x", encoding="utf-8")

    class Raum:
        def sicherstellen(self):
            return wurzel

        def frei(self, pfad):
            return (wurzel in Path(pfad).parents or Path(pfad) == wurzel), "gesperrt"

    monkeypatch.setattr(mobile_routes, "get_dateiraum_service", lambda: Raum())
    assert mobile_routes._file("Notiz.md") == (wurzel / "Notiz.md").resolve()
