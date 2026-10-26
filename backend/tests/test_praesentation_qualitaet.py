from __future__ import annotations

import asyncio

import pytest

from app.services import bildquelle
from app.services.tools import json_retten


def test_kaputtes_folien_json_wird_repariert():
    kaputt = '[{"layout": "title", "title": "Vulkane"} {"layout": "bullets", "title": "Aufbau", "bullets": ["Magma", "Krater",]},]'
    folien = json_retten(kaputt)
    assert [f["layout"] for f in folien] == ["title", "bullets"]
    zeilen = '{"layout": "bullets"\n"title": "Ohne Komma"}'
    assert json_retten(zeilen)["title"] == "Ohne Komma"
    with pytest.raises(ValueError):
        json_retten("das ist kein json")


def test_messgeraete_fotos_kommen_nach_hinten():
    assert bildquelle.nebensache("Photosynthesis System analysing a maize leaf.jpg", "photosynthesis leaf")
    assert not bildquelle.nebensache("Leaf in sunlight.jpg", "photosynthesis leaf")
    assert not bildquelle.nebensache("Geiger counter measurement.jpg", "geiger counter measurement")


def test_bildpruefung_lehnt_ab_und_nimmt_das_naechste(monkeypatch, tmp_path):
    pfade = []
    for name in ("a", "b", "c"):
        pfad = tmp_path / f"{name}.jpg"
        pfad.write_bytes(b"x")
        pfade.append(str(pfad))

    async def geladen(adresse):
        return adresse

    urteile = {pfade[0]: False, pfade[1]: True, pfade[2]: True}

    async def passt(pfad, suche):
        return urteile[pfad]

    monkeypatch.setattr(bildquelle, "_geladen", geladen)
    monkeypatch.setattr(bildquelle, "bild_passt", passt)
    assert asyncio.run(bildquelle.bestes(pfade, "volcano")) == pfade[1]
    urteile[pfade[1]] = False
    urteile[pfade[2]] = False
    assert asyncio.run(bildquelle.bestes(pfade, "volcano")) == ""
    assert asyncio.run(bildquelle.bestes(pfade, "volcano", [0])) == pfade[0]


def test_fehlendes_bild_wird_als_qualitaetshinweis_gemeldet():
    from pptx import Presentation

    from app.services.pptx_service import pruefen

    prs = Presentation()
    for _ in range(3):
        folie = prs.slides.add_slide(prs.slide_layouts[1])
        folie.shapes.title.text = "Titel"
        folie.placeholders[1].text = "Inhalt mit Text"
    folien = [{"layout": "title"}, {"layout": "bullets", "_bild_fehlt": "chloroplast microscope"}, {"layout": "bullets"}]
    hinweise = pruefen(prs, folien)
    assert any("Folie 2" in h and "chloroplast microscope" in h for h in hinweise)
    assert not any("Folie 3:" in h for h in hinweise)


def test_nachtraegliches_bild_verdeckt_keinen_text(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from app.services.pptx_bearbeiten import bearbeiten

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    feld = folie.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(10.5), Inches(4.5))
    feld.text_frame.text = "Lichtintensität: Mehr Licht bedeutet schnellere Photosynthese bis zur Sättigung"
    feld.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    datei = tmp_path / "deck.pptx"
    prs.save(str(datei))
    bild = tmp_path / "bild.png"
    Image.new("RGB", (800, 600), "orange").save(bild)
    ergebnis = asyncio.run(bearbeiten(str(datei), [{"aktion": "bild", "folie": 1, "image": str(bild)}]))
    assert ergebnis.get("ok") or ergebnis.get("erledigt"), ergebnis
    neu = Presentation(str(datei))
    formen = list(neu.slides[0].shapes)
    foto = next(f for f in formen if f.shape_type == 13)
    text = next(f for f in formen if f.has_text_frame and f.text_frame.text.strip())
    assert text.left + text.width <= foto.left
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt < 24


def test_word_mit_tabelle_seitenumbruch_und_fett(tmp_path):
    from docx import Document

    from app.services.dokument_service import DokumentService

    inhalt = "# Bewerbung\n\nIch bin **sehr motiviert**.\n\n---\n\n## Lebenslauf\n\n| Feld | Angabe |\n|---|---|\n| Name | [Dein Name] |\n| Schule | Polytechnische Schule |\n"
    pfad = tmp_path / "test.docx"
    DokumentService._docx(pfad, "Bewerbung", inhalt)
    dokument = Document(str(pfad))
    assert len(dokument.tables) == 1
    tabelle = dokument.tables[0]
    assert tabelle.cell(1, 0).text == "Name" and tabelle.cell(2, 1).text == "Polytechnische Schule"
    assert not any("|" in a.text for a in dokument.paragraphs)
    fett = [l.text for a in dokument.paragraphs for l in a.runs if l.bold]
    assert "sehr motiviert" in fett
    assert any('w:br w:type="page"' in a._p.xml or "w:type=\"page\"" in a._p.xml for a in dokument.paragraphs)


def test_pdf_mit_tabelle(tmp_path):
    pytest.importorskip("reportlab")
    from app.services.dokument_service import DokumentService

    pfad = tmp_path / "test.pdf"
    DokumentService._pdf(pfad, "Test", "## Daten\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n---\n\nSeite zwei")
    assert pfad.stat().st_size > 1000
