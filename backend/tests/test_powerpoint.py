from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest

from app.services.pptx_service import THEMES, get_pptx_service
from app.services.tools import ToolBox, select_tools

pytest.importorskip("pptx")


DECK = [
    {"layout": "title", "title": "Klimawandel", "subtitle": "Ein Ueberblick", "footer": "Lena"},
    {"layout": "bullets", "title": "Ursachen", "bullets": ["CO2: fossile Energie", "Methan: Landwirtschaft"], "notes": "Kurz halten"},
    {"layout": "cards", "title": "Drei Hebel", "items": [{"title": "Strom", "text": "Erneuerbar"}, {"title": "Waerme", "text": "Pumpen"}]},
    {"layout": "stat", "title": "Zahlen", "items": [{"title": "1,5 °C", "text": "Ziel"}, {"title": "420", "text": "ppm CO2"}]},
    {"layout": "two_columns", "title": "Pro und Contra", "items": [{"title": "Pro", "bullets": ["Schnell"]}, {"title": "Contra", "bullets": ["Teuer"]}]},
    {"layout": "timeline", "title": "Fahrplan", "items": [{"title": "2030", "text": "Halbierung"}]},
    {"layout": "quote", "text": "Es gibt keinen Planet B.", "subtitle": "Unbekannt"},
    {"layout": "closing", "title": "Danke", "subtitle": "Fragen?"},
]


def test_praesentation_wird_gebaut(tmp_path):
    ziel = tmp_path / "deck.pptx"
    ergebnis = get_pptx_service().create("Klimawandel", DECK, str(ziel), "forest")
    assert ergebnis["ok"] is True
    assert ziel.exists() and ziel.stat().st_size > 20000
    assert ergebnis["slides"] == len(DECK)
    assert ergebnis["theme"] == "forest"


def test_titelfolie_wird_ergaenzt(tmp_path):
    ziel = tmp_path / "ohne-titel.pptx"
    ergebnis = get_pptx_service().create(
        "Mein Thema", [{"layout": "bullets", "title": "Punkte", "bullets": ["a", "b"]}], str(ziel)
    )
    assert ergebnis["slides"] == 2
    assert ergebnis["layouts"][0] == "title"


def test_unbekanntes_theme_faellt_zurueck(tmp_path):
    ergebnis = get_pptx_service().create("X", DECK[:2], str(tmp_path / "a.pptx"), "quatsch")
    assert ergebnis["theme"] in THEMES


def test_endung_wird_ergaenzt_und_ordner_angelegt(tmp_path):
    ziel = tmp_path / "neu" / "deck"
    ergebnis = get_pptx_service().create("X", DECK[:2], str(ziel))
    assert ergebnis["path"].endswith("deck.pptx")


def test_leere_folienliste_meldet_fehler(tmp_path):
    ergebnis = get_pptx_service().create("X", [], str(tmp_path / "leer.pptx"))
    assert "error" in ergebnis


def test_lesen_liefert_text_und_notizen(tmp_path):
    ziel = tmp_path / "deck.pptx"
    get_pptx_service().create("Klimawandel", DECK, str(ziel))
    gelesen = get_pptx_service().read(str(ziel))
    assert gelesen["folien"] == len(DECK)
    assert any("Ursachen" in t for s in gelesen["inhalt"] for t in s["text"])
    assert any(s["notizen"] for s in gelesen["inhalt"])


def test_werkzeug_erstellt_datei(tmp_path):
    box = ToolBox()
    ziel = tmp_path / "tool.pptx"
    antwort = json.loads(
        asyncio.run(
            box.execute(
                "create_pptx",
                {"title": "Test", "slides": DECK[:3], "path": str(ziel)},
            )
        )
    )
    assert antwort["ok"] is True
    assert ziel.exists()


def test_werkzeug_nimmt_slides_als_json_string(tmp_path):
    box = ToolBox()
    ziel = tmp_path / "string.pptx"
    antwort = json.loads(
        asyncio.run(
            box.execute(
                "create_pptx",
                {"title": "Test", "slides": json.dumps(DECK[:2]), "path": str(ziel)},
            )
        )
    )
    assert antwort["ok"] is True


def test_werkzeug_wird_bei_praesentationsfragen_angeboten():
    assert "create_pptx" in (select_tools("mach mir eine praesentation ueber hunde") or set())
    assert "create_pptx" in (select_tools("bau mir bitte folien fuer mein referat") or set())


def test_relativer_pfad_landet_in_jons_ordner(tmp_path, monkeypatch):
    from app.services import dateiraum_service

    class Raum:
        def zielpfad(self, name, wunsch="", neu=False):
            ordner = tmp_path / "Jon" / (wunsch or "Documents")
            ordner.mkdir(parents=True, exist_ok=True)
            return {"pfad": str(ordner / name)}

    monkeypatch.setattr(dateiraum_service, "get_dateiraum_service", lambda: Raum())
    ergebnis = get_pptx_service().create("Probe", DECK[:2], "Jon-Test.pptx")
    assert ergebnis["ok"]
    assert Path(ergebnis["path"]) == tmp_path / "Jon" / "Documents" / "Jon-Test.pptx"
    assert Path(ergebnis["path"]).is_file()


def test_hell_design_mit_bildobjekten_und_schlussfolie(tmp_path):
    from PIL import Image
    from pptx import Presentation

    from app.services.pptx_service import get_pptx_service

    foto = tmp_path / "foto.jpg"
    Image.new("RGB", (800, 600), (40, 90, 160)).save(foto)
    folien = [
        {"layout": "title", "title": "Künstliche Intelligenz", "subtitle": "Referat", "image": {"path": str(foto), "alt": "KI"}},
        {"layout": "bullets", "title": "Was ist ein Algorithmus?", "bullets": ["Schritt-für-Schritt-Anleitung", "Feste Reihenfolge"], "image": {"pfad": str(foto)}},
        {"layout": "bullets", "title": "Wo gibt es KI?", "bullets": ["Medizin", "Verkehr"], "image": str(foto)},
        {"layout": "closing", "title": "Vielen Dank für Ihre Aufmerksamkeit"},
    ]
    ergebnis = get_pptx_service().create("KI", folien, str(tmp_path / "ki.pptx"))
    assert ergebnis["ok"] and ergebnis["theme"] == "hell"
    assert not any("keinen Inhalt" in h for h in ergebnis["qualitaet"])
    prs = Presentation(ergebnis["path"])
    bilder = [sum(1 for f in folie.shapes if f.shape_type == 13) for folie in prs.slides]
    assert bilder[:3] == [1, 1, 1]
    rund = [f for f in prs.slides[1].shapes if f.shape_type == 13][0]
    assert "ellipse" in rund._element.xml
    assert "Deko-Bogen" in {f.name for f in prs.slides[1].shapes}
    assert "Akzentlinie" in {f.name for f in prs.slides[2].shapes}
    namen = {f.name for f in prs.slides[3].shapes}
    assert "Deko-Ring" in namen
