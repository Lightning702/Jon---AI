from __future__ import annotations

import asyncio

from app.services.aktionen_normal import liste, normalisieren
from app.services.pptx_bearbeiten import AKTION_NAMEN, BILDFELDER


def test_aktionsnamen_und_folienindex_werden_verstanden():
    assert normalisieren({"typ": "bild_hinzufuegen", "slide_index": 2, "bild_suche": "backup"}, AKTION_NAMEN, BILDFELDER, True) == {"typ": "bild_hinzufuegen", "slide_index": 2, "bild_suche": "backup", "aktion": "bild", "folie": 3}
    assert normalisieren({"slide_index": 0, "bild_suche": "cloud"}, AKTION_NAMEN, BILDFELDER, True)["aktion"] == "bild"
    assert normalisieren({"aktion": "uebergang", "art": "fade", "folie": 2}, AKTION_NAMEN, BILDFELDER, True) == {"aktion": "uebergang", "art": "fade", "folie": 2}
    assert normalisieren({"action": "replace_text", "slide": "4"}, AKTION_NAMEN, BILDFELDER, True)["folie"] == 4


def test_liste_versteht_andere_namen():
    assert liste({"slides": [{"a": 1}]}) == [{"a": 1}]
    assert liste({"aktionen": [{"b": 2}], "slides": [{"a": 1}]}) == [{"b": 2}]
    assert liste({}) == []


def test_bilder_kommen_auch_mit_fremden_feldnamen_in_die_folien(tmp_path, monkeypatch):
    from PIL import Image
    from pptx import Presentation

    from app.services import bildquelle
    from app.services.pptx_bearbeiten import bearbeiten
    from app.services.pptx_service import get_pptx_service

    foto = tmp_path / "foto.png"
    Image.new("RGB", (64, 48), (40, 90, 160)).save(foto)

    async def besorgen(eintrag, breit=True):
        return str(foto)

    monkeypatch.setattr(bildquelle, "besorgen", besorgen)
    folien = [{"title": f"Folie {n}", "bullets": ["Punkt"]} for n in range(1, 5)]
    pfad = get_pptx_service().create("Bildtest", folien, str(tmp_path / "bildtest.pptx"), "hell", "", True)["path"]
    vorher = [sum(1 for f in folie.shapes if f.shape_type == 13) for folie in Presentation(pfad).slides]
    ergebnis = asyncio.run(bearbeiten(pfad, [{"typ": "bild_hinzufuegen", "slide_index": 2, "bild_suche": "backup"}, {"slide_index": 3, "bild_suche": "cloud"}]))
    assert not ergebnis.get("error"), ergebnis
    nachher = [sum(1 for f in folie.shapes if f.shape_type == 13) for folie in Presentation(pfad).slides]
    assert nachher[2] > vorher[2] and nachher[3] > vorher[3]
